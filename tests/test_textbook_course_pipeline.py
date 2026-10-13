# -*- coding: utf-8 -*-
"""textbook_course_pipeline 自检：解析器（确定性部分，不触发 TTS/ffmpeg）。

覆盖：章节 sheet 过滤 / 标题 vs 正文判定 / 【网络校准】跳过 / 段落切分与编号。
用法：python -m pytest -q tests/test_textbook_course_pipeline.py
"""
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "scripts"))

import textbook_course_pipeline as tcp  # noqa: E402


def _make_book(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "0封面"
    ws0.append(["《测试教材》"])                    # 前置 sheet → 应排除

    ws1 = wb.create_sheet("1绪论")
    for r in [
        ["章", "节", "小节", "段", "内容"],          # 表头 → 应跳过
        ["01", "00", "00", "00", "第1章 绪论与测试体系"],
        ["01", "00", "00", "00", "【网络校准·2026-08-07】应被跳过"],
        ["", "01", "00", "00", "1.1 学科对象与地位"],
        ["", "", "01", "00", "1.1.1 定义与范畴"],
        ["", "", "", "01", "1.1.1.1 核心概念界定"],
        ["", "", "", "02", "中医基础理论是研究与阐述中医学基本概念、基本原理和基本规律的学科，以气阴阳五行为方法论。"],
        ["", "", "02", "00", "1.1.2 学科位置"],
        ["", "", "", "01", "它在百二十门学科体系中处于核心枢纽地位，上下辐射，是必经节点。"],
    ]:
        ws1.append(r)

    ws2 = wb.create_sheet("2哲思")
    for r in [
        ["章", "节", "小节", "段", "内容"],
        ["02", "00", "00", "00", "第2章 哲学基础"],
        ["", "01", "00", "00", "2.1 气一元论"],
        ["", "", "01", "00", "2.1.1 气的含义"],
        ["", "", "", "01", "气是构成人体和维持生命活动的最基本物质，具有运动不息的特征。"],
    ]:
        ws2.append(r)

    ws3 = wb.create_sheet("ICD11统一命名")          # 附件 sheet → 应排除
    ws3.append(["附"])
    wb.save(path)


@pytest.fixture
def book(tmp_path):
    p = tmp_path / "测试教材.xlsx"
    _make_book(p)
    return p


def test_chapter_sheets_filters_and_orders(book):
    chs = tcp.chapter_sheets(book)
    assert [c[0] for c in chs] == [1, 2]                       # 排除 0* 与附件，按章号升序
    assert [c[1] for c in chs] == ["1绪论", "2哲思"]


@pytest.mark.parametrize("text,expected", [
    ("第1章 绪论与测试体系", True),
    ("1.1 学科对象与地位", True),
    ("1.1.1.1 核心概念界定", True),
    ("麻黄", True),                                            # 短标题
    ("中医基础理论是研究与阐述中医学基本概念、基本原理和基本规律的学科。", False),
])
def test_is_heading(text, expected):
    assert tcp._is_heading(text) is expected


def test_parse_chapter_extracts_prose_and_skips_noise(book):
    segs = tcp.parse_chapter(book, "1绪论")
    assert 2 <= len(segs) <= tcp.MAX_SEGS
    assert [s["no"] for s in segs] == list(range(1, len(segs) + 1))   # 1 起连续编号
    assert all(s["title"] and s["text"] for s in segs)
    joined = " ".join(s["text"] for s in segs)
    assert "中医基础理论是研究" in joined
    assert "百二十门学科" in joined
    assert "网络校准" not in joined                                   # 【…】注记不得混入


def test_parse_chapter_second_chapter(book):
    segs = tcp.parse_chapter(book, "2哲思")
    assert segs and segs[0]["text"].startswith("气是构成人体")


def test_parse_chapter_trailing_empty_column(tmp_path):
    """回归: 表头“内容”列不在行尾(尾部有空列)时仍取到正确内容列。

    实测 2026-10-13: 《方剂学》9 章(解表/清热/固涩/理血/治风/祛痰/消食/治痈疡/AI)
    的 sheet 有 6 列(末尾空列)，旧实现 cells[-1] 读到空 → 0 段、被静默跳过。
    """
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1绪论"
    for r in [
        ["章", "节", "小节", "段", "内容", ""],                 # 6 列，末列为空
        ["01", "", "", "", "第1章 绪论", ""],
        ["", "1", "", "", "1.1 学科对象", ""],
        ["", "", "1", "", "1.1.1 定义", ""],
        ["", "", "", "1", "这是正文，应当被解析为段落文本。", ""],
    ]:
        ws.append(r)
    wb.save(p)
    segs = tcp.parse_chapter(p, "1绪论")
    assert segs, "尾部空列导致 0 段（回归）"
    assert "正文" in segs[0]["text"]
