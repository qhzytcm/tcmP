# -*- coding: utf-8 -*-
"""
gen_yy25_docx.py —— 生成《灵枢·阴阳二十五人》参数化生成式AI模型交付文档 yy25.docx
内容：模型架构图 · 落地路线图 · 男女参数异同 · 模型预测疾病易感性
依据：constitution/ 模块（按九类 Hermes 独到思考重构）
输出：C:\\Users\\DELL\\Desktop\\yy25.docx
"""
import os, sys, json
ROOT = r'C:\Users\DELL\tcmP'
sys.path.insert(0, ROOT)
TMP = os.environ.get('TEMP', r'C:\Users\DELL\AppData\Local\Temp')
DESKTOP = r'C:\Users\DELL\Desktop'

# ── matplotlib 中文字体 ──────────────────────────────────────
import matplotlib
matplotlib.use('Agg')
from matplotlib import font_manager, pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
font_manager.fontManager.addfont(r'C:\Windows\Fonts\simhei.ttf')
plt.rcParams['font.family'] = 'SimHei'
plt.rcParams['axes.unicode_minus'] = False

from constitution.model import ConstitutionModel
m = ConstitutionModel()

WX_COLOR = {'木': '#2E8B57', '火': '#C0392B', '土': '#D4AC0D', '金': '#7F8C8D', '水': '#2C3E50'}
BG = '#F4F7FA'


def _box(ax, x, y, w, h, text, fc='#FFFFFF', ec='#95A5A6', fs=9, tc='#2C3E50', bold=False, lw=1.2):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.006,rounding_size=0.012',
                                linewidth=lw, edgecolor=ec, facecolor=fc, zorder=2))
    ax.text(x + w / 2, y + h / 2, text, ha='center', va='center', fontsize=fs,
            color=tc, zorder=3, fontweight=('bold' if bold else 'normal'), linespacing=1.4)


def _arrow(ax, xy1, xy2, color='#7F8C8D', lw=1.4, style='-|>'):
    ax.add_patch(FancyArrowPatch(xy1, xy2, arrowstyle=style, mutation_scale=12,
                                 linewidth=lw, color=color, zorder=1,
                                 shrinkA=2, shrinkB=2))


# ══════════════════════════════════════════════════════════════
# 图1：模型架构图
# ══════════════════════════════════════════════════════════════
def draw_architecture(png):
    fig, ax = plt.subplots(figsize=(13.6, 8.6))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis('off')

    ax.text(0.5, 0.985, '《灵枢·阴阳二十五人》参数化生成式 AI 模型 — 架构图',
            ha='center', va='top', fontsize=15, fontweight='bold', color='#1A2733')
    ax.text(0.5, 0.955, 'framework_spec v8.0 体质层(Constitution Layer)　|　按九类 Hermes 独到思考重构',
            ha='center', va='top', fontsize=9.5, color='#5D6D7E')

    LX, LW = 0.015, 0.60      # 主列
    RX, RW = 0.635, 0.35      # 右列(九类思考)

    # L0 数据源
    _box(ax, LX, 0.855, LW, 0.075,
         'L0 素材层：桌面四文件 · 黄帝内经全文.xls(LS64/65+SW01) · HDNJ音频理解.xls(讲稿) · 内因七情.xlsx · 外因六淫.xlsx',
         fc='#EAF2FB', ec='#5B9BD5', fs=8.2, bold=True)
    _arrow(ax, (LX + LW / 2, 0.855), (LX + LW / 2, 0.825))

    # L1 体质层 四类参数
    ax.text(LX, 0.822, 'L1 体质层参数 B（四类·高斯归一化 5 区间 L1–L5）', fontsize=9,
            color='#1A2733', fontweight='bold', va='top')
    px = LX; pw = (LW - 0.018) / 4
    labels = [('H 形质', '形(体型五形)\n色(青赤黄白黑)'),
              ('S 性情', '性格倾向\n五常(仁礼信义智)'),
              ('Y 音经', '五音(角徵宫商羽)\n主经/亚型经+情态'),
              ('X 血气承载', 'V_气 / V_血\n须髯六维 / 时令(★性别轴)')]
    for i, (t, d) in enumerate(labels):
        fc = '#FDECEA' if i == 3 else '#FFFFFF'
        _box(ax, px, 0.700, pw, 0.098, f'{t}\n{d}', fc=fc, ec='#C0392B' if i == 3 else '#95A5A6',
             fs=7.8, bold=True)
        px += pw + 0.006
    _arrow(ax, (LX + LW / 2, 0.700), (LX + LW / 2, 0.672))

    # L2 25型矩阵
    ax.text(LX, 0.669, 'L2 二十五型矩阵（五形 × 五亚型 = 5×5）', fontsize=9,
            color='#1A2733', fontweight='bold', va='top')
    px = LX; pw = (LW - 0.024) / 5
    for x in ['木', '火', '土', '金', '水']:
        sub = '/'.join(t['name'] for t in m.types[x])
        _box(ax, px, 0.560, pw, 0.088, f'{x}形\n(主型+4亚型)', fc=WX_COLOR[x], ec=WX_COLOR[x],
             fs=8, tc='white', bold=True)
        ax.text(px + pw / 2, 0.545, yin_of(x), ha='center', va='top', fontsize=7, color=WX_COLOR[x])
        px += pw + 0.006
    _arrow(ax, (LX + LW / 2, 0.560), (LX + LW / 2, 0.532))

    # L3 三大引擎
    ex = LX; ew = (LW - 0.012) / 3
    eng = [('generate() 正向生成', '体质→表型参数向量\n形/色/音/经/情态/血气'),
           ('infer() 逆转推理', '表型观测→25型排序\n(从外知内)'),
           ('predict_susceptibility()', '五形×性别×形色×年忌\n→疾病易感画像')]
    for i, (t, d) in enumerate(eng):
        _box(ax, ex, 0.432, ew, 0.098, f'{t}\n{d}', fc='#EAF7EF', ec='#27AE60', fs=7.6, bold=True)
        ex += ew + 0.006
    _arrow(ax, (LX + LW / 2, 0.432), (LX + LW / 2, 0.404))

    # L4 输出层
    ox = LX; ow = (LW - 0.012) / 3
    outs = [('疾病易感画像', '主脏/经络\n易感病谱/高危窗'),
            ('个体化调治', '五音五味·五谷畜果\n刺约(个体化针刺)'),
            ('性别参数对照', '男(血气均衡·须髯)\n女(气余血少·无须)')]
    for i, (t, d) in enumerate(outs):
        _box(ax, ox, 0.304, ow, 0.098, f'{t}\n{d}', fc='#FFF8E1', ec='#D4AC0D', fs=7.6, bold=True)
        ox += ow + 0.006
    _arrow(ax, (LX + LW / 2, 0.304), (LX + LW / 2, 0.276))

    # L5 耦合层
    _box(ax, LX, 0.190, LW, 0.080,
         'L5 与病因层耦合：Φ_total = ∫ Σ [ B_体质增益因子 ] · R_j(E,P; Z_j) dτ - D0\n'
         '体质层(宿主禀赋) × 病因层(内因七情 / 外因六淫 / 不内外因)',
         fc='#EFE7F7', ec='#8E44AD', fs=8.2, bold=True)
    _arrow(ax, (LX + LW / 2, 0.190), (LX + LW / 2, 0.162))

    # L6 数字人
    _box(ax, LX, 0.082, LW, 0.075,
         'sage-api 数字人接入：/constitution 端点 · 25型问诊对话流 · 六者终端画像',
         fc='#EAF2FB', ec='#5B9BD5', fs=8.2, bold=True)

    # ── 右列：九类 Hermes 独到思考 ──
    _box(ax, RX, 0.915, RW, 0.055, '九类 Hermes 独到思考 → 现代对话', fc='#1A2733', ec='#1A2733',
         fs=9.5, tc='white', bold=True)
    y = 0.888; bh = 0.0825
    for h in m.hermes9():
        yy = y - bh
        _box(ax, RX, yy, RW, bh - 0.006,
             f'{h["no"]}·{h["title"].split("=")[0]}\n→ {h["modern"]}',
             fc='#F7FAFC', ec='#B0BEC5', fs=6.6, tc='#37474F')
        y = yy

    fig.savefig(png, dpi=185, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return png


def yin_of(x):
    return {'木': '角', '火': '徵', '土': '宫', '金': '商', '水': '羽'}[x]


# ══════════════════════════════════════════════════════════════
# 图2：落地路线图
# ══════════════════════════════════════════════════════════════
def draw_roadmap(png):
    fig, ax = plt.subplots(figsize=(13.6, 4.6))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis('off')
    ax.text(0.5, 0.95, '阴阳二十五人模型 — 落地路线图（P1→P8）',
            ha='center', va='top', fontsize=14, fontweight='bold', color='#1A2733')

    rp = m.roadmap()
    SC = {'已完成': '#27AE60', '待接线': '#F39C12', '待开发': '#E67E22', '规划中': '#95A5A6'}
    n = len(rp)
    x0, gap = 0.02, 0.0035
    bw = (1 - 2 * x0 - gap * (n - 1)) / n
    # 时间轴
    ax.plot([x0, 1 - x0], [0.40, 0.40], color='#B0BEC5', lw=2, zorder=1)
    for i, p in enumerate(rp):
        x = x0 + i * (bw + gap)
        c = SC.get(p['status'], '#95A5A6')
        # 节点
        ax.add_patch(plt.Circle((x + bw / 2, 0.40), 0.012, color=c, zorder=3))
        # 上方阶段名
        _box(ax, x, 0.46, bw, 0.20, p['phase'], fc=c, ec=c, fs=8, tc='white', bold=True)
        # 下方交付
        _box(ax, x, 0.02, bw, 0.33, f'{p["deliver"]}\n\n依赖: {p["dep"]}', fc='#F7FAFC', ec='#CFD8DC', fs=6.3)
        # 连接线
        _arrow(ax, (x + bw / 2, 0.46), (x + bw / 2, 0.412), color='#B0BEC5', lw=1)
        _arrow(ax, (x + bw / 2, 0.388), (x + bw / 2, 0.35), color='#B0BEC5', lw=1)
        ax.text(x + bw / 2, 0.685, p['status'], ha='center', va='bottom', fontsize=7, color=c,
                fontweight='bold')
    # 图例
    lg = '　'.join(f'{k}' for k in SC)
    ax.text(0.5, 0.905, '状态：' + lg, ha='center', va='top', fontsize=8, color='#5D6D7E')
    fig.savefig(png, dpi=185, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return png


# ══════════════════════════════════════════════════════════════
# docx 组装
# ══════════════════════════════════════════════════════════════
from docx import Document
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement


def set_style_font(style, name='微软雅黑', size=10.5):
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), name)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell(cell, text, size=9, bold=False, color=None, align='left'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    for i, line in enumerate(str(text).split('\n')):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = '微软雅黑'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        if color:
            r.font.color.rgb = RGBColor(*color)


def make_table(doc, headers, rows, widths=None, hsize=9, bsize=8.4, hfill='DCE6F1'):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, size=hsize, bold=True, align='center')
        shade(t.rows[0].cells[j], hfill)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            set_cell(cells[j], v, size=bsize, align='center' if j == 0 else 'left')
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Mm(w)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def body(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p


def build_docx(png_arch, png_road, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin = sec.bottom_margin = Mm(18)
    set_style_font(doc.styles['Normal'], '微软雅黑', 10.5)
    for nm, sz in [('Title', 19), ('Heading 1', 15), ('Heading 2', 12.5)]:
        if nm in [s.name for s in doc.styles]:
            set_style_font(doc.styles[nm], '微软雅黑', sz)

    # 封面
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('《灵枢·阴阳二十五人》参数化生成式 AI 模型'); r.font.size = Pt(21); r.font.bold = True
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run('—— 按九类 Hermes 独到思考重构 · 男女异同参数 · 疾病易感预测 ——')
    r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x5D, 0x6D, 0x7E)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    m1 = doc.add_paragraph(); m1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m1.add_run('篇目：灵枢第六十四篇 阴阳二十五人（配 灵枢·五音五味 / 素问·上古天真论）\n'
                   '契约：framework_spec v8.0 体质层　|　实现：tcmP/constitution/　|　文件：yy25.docx')
    r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 摘要
    doc.add_heading('摘要', level=1)
    body(doc, '本模型将《灵枢·阴阳二十五人》“五五二十五人”体质-人格分类形式化为参数化生成式模型：'
              '以 木火土金水 五形 × 五种变异的 25 型为分类骨架，以「形态、性格、气候适应、经络、五音」五维刻画，'
              '落地为四类参数 B=(形质 H, 性情 S, 音经 Y, 血气承载 X)，并用与病因层同款的'
              '「高斯归一化 + 5 区间 L1–L5」标尺统一度量。模型含三大引擎——正向生成 generate()、'
              '逆转推理 infer()（从外知内）、疾病易感预测 predict_susceptibility()；'
              '作为「体质层（宿主禀赋）」与病因层（内因七情/外因六淫/不内外因）经泛函相乘耦合。'
              '本文按九类 Hermes 独到思考重构，并给出模型架构图、落地路线图、男女参数异同对照与疾病易感预测。')

    # 一、模型架构图
    doc.add_page_break()
    doc.add_heading('一、模型架构图', level=1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png_arch, width=Inches(6.75))
    caption(doc, '图1　阴阳二十五人参数化生成式 AI 模型架构（体质层 B 参数 → 25型矩阵 → 三大引擎 → 输出/耦合/数字人；右列为九类 Hermes 思考）')

    doc.add_heading('1.1 四类参数体系（模型参数空间 B）', level=2)
    make_table(doc,
               ['参数类', '符号维度', '取值', '作用'],
               [['H 形质参数', 'shape / se', '体型五形特征向量；五色 青赤黄白黑', '由外知内的外貌锚（形色相得判据输入）'],
                ['S 性情参数', 'xingqing / chang', '才/劳心/多虑/轻财/静悍…；五常 仁礼信义智', '人格心理轴（五形人格→疾病关联）'],
                ['Y 音经参数', 'yin / jing / tai', '角徵宫商羽；主经+亚型经；情态 佗佗然…', '体质-脏器功能轴 + 亚型分化'],
                ['X 血气承载（★性别轴）', 'V_qi / V_xue / hair', '0–1（5区间）；须髯六维；时令适应', '性别差异核心：男有须髯、女无须']],
               widths=[30, 30, 60, 45])

    doc.add_heading('1.2 九类 Hermes 独到思考（模型理论骨架）', level=2)
    rows = [[str(h['no']), h['title'], h['core'], h['modern']] for h in m.hermes9()]
    make_table(doc, ['#', '独到思考', '中医内核', '现代对话'], rows,
               widths=[8, 42, 62, 53], hsize=8.5, bsize=7.6)

    # 二、落地路线图
    doc.add_page_break()
    doc.add_heading('二、落地路线图', level=1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png_road, width=Inches(6.75))
    caption(doc, '图2　落地路线图 P1→P8（绿=已完成 / 黄橙=待接线开发 / 灰=规划中）')
    make_table(doc, ['阶段', '交付物', '依赖', '状态'],
               [[r['phase'], r['deliver'], r['dep'], r['status']] for r in m.roadmap()],
               widths=[32, 78, 35, 20], hsize=9, bsize=8)

    # 三、男女参数异同
    doc.add_page_break()
    doc.add_heading('三、男女参数异同', level=1)
    cg = m.compare_gender()
    body(doc, '依据《灵枢·五音五味》“妇人无须者……冲、任之脉，不荣口唇，故须不生焉；今妇人之生，'
              '有余于气，不足于血，以其数脱血也”，及《素问·上古天真论》“女子七岁……丈夫八岁……”。',
         size=9.5, color=(0x5D, 0x6D, 0x7E))
    doc.add_heading('3.1 相同点（男 = 女，9 项）', level=2)
    make_table(doc, ['维度', '参数', '男', '女', '依据'],
               [[r['dim'], r['param'], r['male'], r['female'], r['src']] for r in cg['same']],
               widths=[22, 30, 48, 48, 20], hsize=8.5, bsize=7.6)
    doc.add_heading('3.2 不同点（男 ≠ 女，8 项）', level=2)
    make_table(doc, ['维度', '参数', '男', '女', '依据'],
               [[r['dim'], r['param'], r['male'], r['female'], r['src']] for r in cg['diff']],
               widths=[22, 30, 48, 48, 20], hsize=8.5, bsize=7.6)
    doc.add_heading('3.3 结论', level=2)
    body(doc, '性别差异不在“分型骨架”，而在“血气承载轴 X”：① 女性 V_血 系统性偏低、V_气 相对偏高（数脱血）；'
              '② 由此派生“须髯维度失效”，须以“冲任荣唇/月事”轴替代；'
              '③ 叠加“女七男八”生殖节律与特殊子类（男：宦者/天宦；女：妇人/五不女）。'
              '25 型的五行骨架、五色五音、五脏主经、性情五常、时令适应、年忌与调治总则男女通用。')

    # 四、模型预测疾病易感性
    doc.add_page_break()
    doc.add_heading('四、模型预测疾病易感性', level=1)
    body(doc, '模型以“五形主脏 → 病机 → 易感病谱”，叠加“性别修饰”“形色相得/相克”“年忌高危窗”，'
              '输出个体疾病易感画像。其中火形（A型人格）→ 心血管病，对应现代行为心脏病学，预测效力最强。', size=9.5)
    rows = []
    for s in m.susceptibility_all():
        rows.append([f'{s["xing"]}形', f'{s["zang"]}({s["jing"]})', s['core_patho'],
                     '、'.join(s['diseases']), s['modern'], s['sex_mod'], s['high_window']])
    make_table(doc, ['五形', '主脏(经)', '核心病机', '易感病谱', '现代对应', '性别差异修饰', '高危窗'],
               rows, widths=[12, 16, 34, 44, 30, 34, 20], hsize=8.2, bsize=7.2)
    doc.add_heading('4.1 易感规则与算例', level=2)
    body(doc, f'规则：① 形色相得/相生 → 易感基线；形胜色/色胜形（相克）→ 易感整体上浮。'
              f'② 年忌（{", ".join(map(str, [7,16,25,34,43,52,61]))} 岁）为高敏感窗。'
              f'③ 性别修饰：女（气余血少）偏情志/血虚/经断前后；男（血气均衡）偏代谢/心血管早发。', size=9.5)
    ex = m.predict_susceptibility('火', None, 'female')
    body(doc, f'算例（火形·女）：主脏 {ex["zang"]}({ex["jing"]})，病机 {ex["core_patho"]}；'
              f'形色关系 {ex["form_color"]} → 风险 {ex["risk_level"]}；'
              f'易感：{"、".join(ex["diseases"])}；性别修饰：{ex["sex_specific"]}；高危窗：{ex["high_window"]}。', size=9.5)

    # 附录：25型编码表
    doc.add_page_break()
    doc.add_heading('附录：二十五型编码表', level=1)
    rows = []
    for x in ['木', '火', '土', '金', '水']:
        subs = m.types[x]
        rows.append([x + '形', yin_of(x), m.xing[x]['zhu_jing_cn'], m.xing[x]['se'],
                     subs[0]['name'] + '(主型)', '、'.join(t['name'] for t in subs[1:])])
    make_table(doc, ['五行', '五音', '主经', '色', '主型', '四亚型'],
               rows, widths=[16, 14, 34, 12, 30, 62], hsize=9, bsize=8.2)
    body(doc, '注：《灵枢·阴阳二十五人》与《灵枢·五音五味》对火形第五型名称（质判/判徵）略有出入，'
              '本表并存标注。阴阳之人（太阴/少阴/太阳/少阳/阴阳和平，《灵枢·通天》）为 25 型之外的并列分类。',
         size=9, color=(0x80, 0x80, 0x80))

    doc.save(out)
    return out


if __name__ == '__main__':
    png_arch = os.path.join(TMP, 'yy25_arch.png')
    png_road = os.path.join(TMP, 'yy25_roadmap.png')
    draw_architecture(png_arch)
    print('架构图:', png_arch, os.path.getsize(png_arch))
    draw_roadmap(png_road)
    print('路线图:', png_road, os.path.getsize(png_road))
    out = os.path.join(DESKTOP, 'yy25.docx')
    build_docx(png_arch, png_road, out)
    print('docx:', out, os.path.getsize(out))
