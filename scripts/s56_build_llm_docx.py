# -*- coding: utf-8 -*-
"""
S56 用 LLM 矫正结果重建 docx（03 底本）
读 data/pdf3_units_v2.json + data/pdf3_llm_out.jsonl（逐单元矫正结果）
未矫正/护栏退回者保留原文并标注。输出：
  C:\\Users\\DELL\\Desktop\\03-bycd-LLM校本文.docx
  C:\\Users\\DELL\\Desktop\\03-bycd-LLM校本文-三段结构.docx
"""
import os, json, sys
from collections import Counter
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
UNITS = os.path.join(ROOT, 'data', 'pdf3_units_v2.json')
OUTL = os.path.join(ROOT, 'data', 'pdf3_llm_out.jsonl')
DESKTOP = r'C:\Users\DELL\Desktop'
ORD = ('病源', '病状', '治法')


def load_out():
    m = {}
    if os.path.exists(OUTL):
        for line in open(OUTL, encoding='utf-8'):
            try:
                r = json.loads(line)
                m[r['i']] = r
            except Exception:
                pass
    return m


def secs(b):
    pos = {}
    for k in ORD:
        i = b.find(k)
        if i >= 0:
            pos[k] = i
    if not pos:
        return {'全': b}
    order = sorted(pos.items(), key=lambda kv: kv[1])
    out = {}
    for i, (k, p) in enumerate(order):
        e = order[i + 1][1] if i + 1 < len(order) else len(b)
        s = b[p:e]
        if s.startswith(k):
            s = s[len(k):]
        out[k] = s
    if order[0][1] > 0:
        out[order[0][0]] = b[:order[0][1]] + out[order[0][0]]
    return out


def main():
    units = json.load(open(UNITS, encoding='utf-8'))
    outs = load_out()
    body = []
    st = Counter()
    for i, u in enumerate(units):
        r = outs.get(i)
        txt = (r['out'] if (r and r.get('ok')) else u['body']) if r else u['body']
        # 分类标注：正常校 / 护栏退回 / 过短 / **瞬时异常未校（待重试）** / 未校
        if r and r.get('ok'):
            flag = 'LLM校'
        elif not r:
            flag = '未校'
        elif r.get('why') == 'too_short':
            flag = '过短未校'
        elif r.get('why'):
            flag = f"未校({r['why']})"
        else:
            flag = '护栏退回'
        s = secs(txt)
        for k in ORD:
            if k in s:
                st[k] += 1
        if all(k in s for k in ORD):
            st['三段齐备'] += 1
        body.append((u['name'], txt, s, flag))
    n_llm = sum(1 for x in body if x[3] == 'LLM校')
    print(f'单元 {len(units)} | LLM已校 {n_llm} | 护栏退回 {sum(1 for x in body if x[3]=="护栏退回")} | '
          f'未校 {sum(1 for x in body if x[3]=="未校")}')
    print('分段:', dict(st))

    for tag, structured in (('03-bycd-LLM校本文.docx', False),
                            ('03-bycd-LLM校本文-三段结构.docx', True)):
        doc = Document()
        s0 = doc.sections[0]; s0.page_width, s0.page_height = Mm(210), Mm(297)
        stl = doc.styles['Normal']; stl.font.name = '仿宋'; stl.font.size = Pt(10.5)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run('《病源辞典》03 清晰底本·LLM 语义校本'); rr.bold = True; rr.font.size = Pt(15)
        q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run(f'共 {len(body)} 单元　LLM 已校 {n_llm}　三段齐备 {st["三段齐备"]}　'
                  f'（qwen2.5:7b · 护栏：长度比与新增汉字率）').font.size = Pt(9)
        for i, (nm, txt, s, flag) in enumerate(body, 1):
            h = doc.add_paragraph(); r2 = h.add_run(f'{i:04d}　{nm}　〔{flag}〕')
            r2.bold = True; r2.font.size = Pt(11.5)
            if structured:
                for k in ORD:
                    if k in s:
                        doc.add_paragraph(f'{k}：{s[k]}')
                for k, v in s.items():
                    if k not in ORD:
                        doc.add_paragraph(v)
            else:
                doc.add_paragraph(txt)
        out = os.path.join(DESKTOP, tag)
        try:
            doc.save(out)
        except PermissionError:
            stem, ext = os.path.splitext(tag)
            out = os.path.join(DESKTOP, f'{stem}_新{ext}')   # 目标被 Word 占用 → 落盘回退
            doc.save(out)
            print('  [!] 目标被占用（Word 打开中），已改存:', out)
        print('→', out, os.path.getsize(out), 'bytes')


if __name__ == '__main__':
    main()
