# -*- coding: utf-8 -*-
"""
gen_yy25v2_docx.py —— 生成 yy25V2.docx（模型 README V2 的 Word 交付版）
内容：阴阳五行双层进制 · 60甲子终身模式 · 男8女7发育 · 左右镜像 · 双预测目标(健康/证候易感)
依据：constitution/ 模块（V2 术数层）
输出：C:\\Users\\DELL\\Desktop\\yy25V2.docx
"""
import os, sys
ROOT = r'C:\Users\DELL\tcmP'
sys.path.insert(0, ROOT)
TMP = os.environ.get('TEMP', r'C:\Users\DELL\AppData\Local\Temp')
DESKTOP = r'C:\Users\DELL\Desktop'

import matplotlib
matplotlib.use('Agg')
from matplotlib import font_manager, pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
font_manager.fontManager.addfont(r'C:\Windows\Fonts\simhei.ttf')
plt.rcParams['font.family'] = 'SimHei'
plt.rcParams['axes.unicode_minus'] = False

from constitution import framework as F
from constitution.model import ConstitutionModel
m = ConstitutionModel()

WX = {'木': '#2E8B57', '火': '#C0392B', '土': '#D4AC0D', '金': '#7F8C8D', '水': '#2C3E50'}


def _box(ax, x, y, w, h, text, fc='#FFFFFF', ec='#95A5A6', fs=9, tc='#2C3E50', bold=False, lw=1.2):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.006,rounding_size=0.012',
                                linewidth=lw, edgecolor=ec, facecolor=fc, zorder=2))
    ax.text(x + w / 2, y + h / 2, text, ha='center', va='center', fontsize=fs,
            color=tc, zorder=3, fontweight=('bold' if bold else 'normal'), linespacing=1.4)


def _arrow(ax, xy1, xy2, color='#7F8C8D', lw=1.4):
    ax.add_patch(FancyArrowPatch(xy1, xy2, arrowstyle='-|>', mutation_scale=12,
                                 linewidth=lw, color=color, zorder=1, shrinkA=2, shrinkB=2))


# ══════════════════════════════════════════════════════════════
# 图1：V2 模型架构图
# ══════════════════════════════════════════════════════════════
def draw_arch(png):
    fig, ax = plt.subplots(figsize=(12.8, 8.4))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis('off')
    ax.text(0.5, 0.985, 'yy25 V2 — 体质-术数生成式模型 架构图', ha='center', va='top',
            fontsize=15.5, fontweight='bold', color='#1A2733')
    ax.text(0.5, 0.952, '阴阳五行双层进制 · 60甲子终身模式 · 男8女7发育 · 左右镜像 · 双预测目标',
            ha='center', va='top', fontsize=9.5, color='#5D6D7E')

    LX, LW = 0.02, 0.96
    # L0
    _box(ax, LX, 0.862, LW, 0.062,
         'L0 素材层：黄帝内经 LS64阴阳二十五人 / LS65五音五味 / SW01上古天真论(女七男八)  +  桌面四文件',
         fc='#EAF2FB', ec='#5B9BD5', fs=8.4, bold=True)
    _arrow(ax, (0.5, 0.862), (0.5, 0.83))

    # 术数层
    ax.add_patch(FancyBboxPatch((LX, 0.585), LW, 0.245, boxstyle='round,pad=0.006,rounding_size=0.012',
                                linewidth=1.5, edgecolor='#8E44AD', facecolor='#F6F0FB', zorder=1))
    ax.text(LX + 0.012, 0.822, '术数层（V2 矫正轴）：阴阳五行双层进制 → 60甲子 / 男8女7 / 左右镜像',
            fontsize=9.4, color='#6C3483', fontweight='bold', va='top')
    _box(ax, LX + 0.01, 0.735, 0.225, 0.062, '阴阳五行双层进制\n阴阳(2) × 五行(5) = 10 天干', fc='#EFE7F7', ec='#8E44AD', fs=7.8, bold=True)
    _box(ax, LX + 0.245, 0.735, 0.225, 0.062, '天干(10) × 地支(12)\n最小公倍数 = 60 甲子', fc='#EFE7F7', ec='#8E44AD', fs=7.8, bold=True)
    _box(ax, LX + 0.48, 0.735, 0.225, 0.062, '60甲子 · 终身模式\n本命天干五行 = 命局主行', fc='#EFE7F7', ec='#8E44AD', fs=7.8, bold=True)
    _box(ax, LX + 0.715, 0.735, 0.235, 0.062, '男8女7 · 发育节律\n女七/男八 · 天癸节点', fc='#EFE7F7', ec='#8E44AD', fs=7.8, bold=True)
    _box(ax, LX + 0.01, 0.605, 0.44, 0.105,
         '左右镜像：男 左=阳·右=阴  |  女 右=阳·左=阴\n（亚型方位阴阳 / 经脉取穴 / 脉诊随性别翻转）',
         fc='#FFFFFF', ec='#8E44AD', fs=7.8)
    _box(ax, LX + 0.47, 0.605, 0.48, 0.105,
         '矫正运算：命局五行 × 禀赋五形 → 生/克/同\n发育岁运 · 侧别阴阳 → 叠加修正 25型参数',
         fc='#FFFFFF', ec='#8E44AD', fs=7.8)
    _arrow(ax, (0.5, 0.585), (0.5, 0.553))

    # 体质层
    ax.add_patch(FancyBboxPatch((LX, 0.415), LW, 0.138, boxstyle='round,pad=0.006,rounding_size=0.012',
                                linewidth=1.4, edgecolor='#27AE60', facecolor='#EDF9F1', zorder=1))
    ax.text(LX + 0.012, 0.545, '体质层（V1 骨架）：四类参数 B = (形质H · 性情S · 音经Y · 血气承载X) → 五形×五亚型=25型',
            fontsize=9, color='#1E8449', fontweight='bold', va='top')
    px, pw = LX + 0.01, (LW - 0.06) / 5
    for x in ['木', '火', '土', '金', '水']:
        _box(ax, px, 0.428, pw, 0.062, f'{x}形({ {"木":"角","火":"徵","土":"宫","金":"商","水":"羽"}[x] })\n主型+4亚型',
             fc=WX[x], ec=WX[x], fs=7.6, tc='white', bold=True)
        px += pw + 0.008
    _arrow(ax, (0.5, 0.415), (0.5, 0.383))

    # 预测引擎
    _box(ax, LX + 0.06, 0.268, 0.40, 0.108,
         '目标① 个体健康状态\nhealth_state() → 0–1 健康评分 + 状态分级\n(禀赋⊕命局⊕发育期⊕形色)',
         fc='#FFF8E1', ec='#D4AC0D', fs=8, bold=True)
    _box(ax, LX + 0.50, 0.268, 0.40, 0.108,
         '目标② 疾病证候易感性\nsyndrome_susceptibility() → 证候易感度排序\n(五形⊕命局⊕发育期⊕性别)',
         fc='#FDECEA', ec='#C0392B', fs=8, bold=True)
    _arrow(ax, (0.28, 0.268), (0.28, 0.236))
    _arrow(ax, (0.72, 0.268), (0.72, 0.236))

    # 辅助引擎 + 耦合
    _box(ax, LX + 0.06, 0.150, 0.40, 0.082,
         '辅助引擎：generate() 正向 / infer() 逆推\npredict_susceptibility() 病谱', fc='#F4F6F7', ec='#95A5A6', fs=7.6)
    _box(ax, LX + 0.50, 0.150, 0.40, 0.082,
         '耦合：Φ_total = ∫Σ[B_体质增益]·R_j(E,P;Z_j)dτ - D0\n体质层 × 病因层(七情/六淫/不内外因)', fc='#EFE7F7', ec='#8E44AD', fs=7.4, bold=True)
    _arrow(ax, (0.5, 0.150), (0.5, 0.118), color='#8E44AD')
    _box(ax, LX, 0.055, LW, 0.058, 'sage-api 数字人 /constitution 端点 · 25型画像 · 六者终端',
         fc='#EAF2FB', ec='#5B9BD5', fs=8, bold=True)

    fig.savefig(png, dpi=185, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return png


# ══════════════════════════════════════════════════════════════
# 图2：阴阳五行双层进制 → 60甲子 / 男8女7
# ══════════════════════════════════════════════════════════════
def draw_radix(png):
    fig, ax = plt.subplots(figsize=(13.2, 6.6))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis('off')
    ax.text(0.5, 0.975, '阴阳五行双层进制 → 天干(10) → 60甲子 → 男8女7', ha='center', va='top',
            fontsize=14, fontweight='bold', color='#1A2733')

    # —— A. 阴阳(2) ——
    _box(ax, 0.02, 0.76, 0.14, 0.14, '阴阳\n(2进制)\n阴 · 阳', fc='#EAF2FB', ec='#5B9BD5', fs=8.5, bold=True)
    ax.text(0.17, 0.83, '×', fontsize=16, ha='center', va='center')
    # —— B. 五行(5) ——
    _box(ax, 0.19, 0.76, 0.30, 0.14, '五行 (5进制)\n木  火  土  金  水', fc='#EDF9F1', ec='#27AE60', fs=8.5, bold=True)
    _arrow(ax, (0.50, 0.83), (0.55, 0.83))
    # —— C. 天干(10) 2×5 ——
    ax.text(0.79, 0.915, '天干 = 阴阳 × 五行（2×5=10）', ha='center', fontsize=8.6, color='#6C3483', fontweight='bold')
    wx5 = ['木', '火', '土', '金', '水']
    # 阳干行
    for j, (tg, wx) in enumerate(zip(['甲', '丙', '戊', '庚', '壬'], wx5)):
        _box(ax, 0.56 + j * 0.086, 0.80, 0.078, 0.055, f'{tg}(阳)', fc=WX[wx], ec=WX[wx], fs=8, tc='white', bold=True)
    # 阴干行
    for j, (tg, wx) in enumerate(zip(['乙', '丁', '己', '辛', '癸'], wx5)):
        _box(ax, 0.56 + j * 0.086, 0.735, 0.078, 0.055, f'{tg}(阴)', fc='white', ec=WX[wx], fs=8, tc=WX[wx], bold=True)
    # 五行列注
    for j, wx in enumerate(wx5):
        ax.text(0.56 + j * 0.086 + 0.039, 0.722, wx, ha='center', va='top', fontsize=7.5, color=WX[wx])

    # —— D. 地支(12) ——
    _box(ax, 0.02, 0.55, 0.10, 0.09, '地支\n(12)', fc='#FEF9E7', ec='#D4AC0D', fs=8, bold=True)
    dz_txt = '　'.join(F.DIZHI)
    ax.text(0.13, 0.595, dz_txt, fontsize=9, va='center', color='#7D6608')
    ax.text(0.13, 0.545, '子(阳)丑(阴)相间 · 各含五行与生肖', fontsize=7.5, va='center', color='#999')
    ax.text(0.62, 0.595, '天干(10) × 地支(12)  最小公倍数 = 60', fontsize=9, va='center',
            color='#6C3483', fontweight='bold')

    # —— E. 60甲子 梯列 ——
    ax.text(0.02, 0.49, '60甲子（一甲子 60 年）', fontsize=9.5, color='#1A2733', fontweight='bold', va='top')
    cols, cw = 30, (0.96 - 0.02) / 30
    for i in range(60):
        c, r = i % 30, i // 30
        x = 0.02 + c * cw
        y = 0.40 - r * 0.085
        col = '#8E44AD' if i % 5 == 0 else '#D2B4DE'
        ax.add_patch(FancyBboxPatch((x, y), cw * 0.92, 0.052, boxstyle='round,pad=0.001,rounding_size=0.004',
                                    fc='#F6F0FB', ec=col, lw=0.8, zorder=2))
        ax.text(x + cw * 0.46, y + 0.026, F.JIAZI60[i]['gz'], ha='center', va='center',
                fontsize=(6.4 if i % 5 == 0 else 5.4), color=('#6C3483' if i % 5 == 0 else '#A569BD'))

    # —— F. 男8女7 时间轴 ——
    ax.text(0.02, 0.24, '男8女7 · 发育阶段性模式', fontsize=9.5, color='#1A2733', fontweight='bold', va='top')
    ax.plot([0.03, 0.97], [0.15, 0.15], color='#B0BEC5', lw=2)
    fs = F.DEV_STAGES['female']; ms = F.DEV_STAGES['male']
    for s in fs:
        x = 0.03 + (s['age'] / 64) * 0.94
        ax.plot([x], [0.175], marker='o', color='#C0392B', ms=5, zorder=3)
        ax.text(x, 0.20, f'{s["age"]}', ha='center', fontsize=6.5, color='#C0392B')
    for s in ms:
        x = 0.03 + (s['age'] / 64) * 0.94
        ax.plot([x], [0.125], marker='s', color='#2C3E50', ms=5, zorder=3)
        ax.text(x, 0.085, f'{s["age"]}', ha='center', fontsize=6.5, color='#2C3E50')
    ax.text(0.985, 0.185, '女(7)', ha='right', fontsize=8, color='#C0392B')
    ax.text(0.985, 0.075, '男(8)', ha='right', fontsize=8, color='#2C3E50')
    ax.text(0.03, 0.03, '女：7/14/21/28/35/42/49（二七天癸至·七七竭）    男：8/16/24/32/40/48/56/64（二八至·八八竭）',
            fontsize=7.6, color='#5D6D7E')

    fig.savefig(png, dpi=185, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return png


# ══════════════════════════════════════════════════════════════
# docx 组装
# ══════════════════════════════════════════════════════════════
from docx import Document
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_style_font(style, name='微软雅黑', size=10.5):
    style.font.name = name; style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), name)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell(cell, text, size=9, bold=False, color=None, align='left'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    for i, line in enumerate(str(text).split('\n')):
        if i:
            p.add_run().add_break()
        r = p.add_run(line); r.font.size = Pt(size); r.font.bold = bold
        r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        if color:
            r.font.color.rgb = RGBColor(*color)


def make_table(doc, headers, rows, widths=None, hsize=9, bsize=8.2, hfill='DCE6F1'):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, size=hsize, bold=True, align='center'); shade(t.rows[0].cells[j], hfill)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            set_cell(cells[j], v, size=bsize, align='center' if j == 0 else 'left')
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Mm(w)
    return t


def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def body(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p


def build(png_arch, png_radix, out):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(18); s.top_margin = s.bottom_margin = Mm(18)
    set_style_font(doc.styles['Normal'], '微软雅黑', 10.5)
    for nm, sz in [('Title', 19), ('Heading 1', 15), ('Heading 2', 12.5)]:
        if nm in [st.name for st in doc.styles]:
            set_style_font(doc.styles[nm], '微软雅黑', sz)

    # 封面
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('yy25 V2　体质-术数生成式模型'); r.font.size = Pt(22); r.font.bold = True
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run('阴阳五行双层进制 · 60甲子终身模式 · 男8女7发育 · 左右镜像 · 双预测目标')
    r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x5D, 0x6D, 0x7E)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    mp = doc.add_paragraph(); mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = mp.add_run('《灵枢·阴阳二十五人》体质分类 × 中文术数模式 ｜ 预测目标：①个体健康状态 ②疾病证候易感性\n'
                   '契约 framework_spec v8.0 体质层 ｜ 实现 tcmP/constitution/ ｜ 文件 yy25V2.docx')
    r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 摘要
    doc.add_heading('摘要', level=1)
    body(doc, 'yy25 V2 在 V1（五形×五亚型 25 型、性别参数）基础上，引入**中文术数模式**作为矫正思路：'
              '以「阴阳(2进制) × 五行(5进制) = 天干(10)」为复合双层进制，再以「天干(10) × 地支(12) = 60甲子」'
              '生成个体终身模式；以「男8女7」生成性别发育阶段性模式；以「男左为阳右为阴、女右为阳左为阴」的'
              '镜像规则校正侧别阴阳。模型据此输出两大预测目标——① 个体健康状态（0–1 健康评分与分级）、'
              '② 疾病证候易感性（证候易感度排序）。')

    # 一、理论基石：双层进制
    doc.add_page_break()
    doc.add_heading('一、理论基石：阴阳五行双层进制', level=1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png_radix, width=Inches(6.75))
    caption(doc, '图1　阴阳五行双层进制 → 天干(10) → 60甲子 → 男8女7 发育节律')
    make_table(doc, ['进制', '基数', '位元', '含义'],
               [['低位', '2（阴阳）', '阳 / 阴', '属性两极（二进制）'],
                ['高位', '5（五行）', '木 火 土 金 水', '属性五分（五进制）'],
                ['复合', '2 × 5 = 10', '天干', '甲丙戊庚壬(阳) · 乙丁己辛癸(阴)'],
                ['再叠加', '10 × 12', '60甲子', '天干(10) 与 地支(12) 最小公倍数 = 60']],
               widths=[20, 28, 34, 88])

    # 二、模型架构
    doc.add_page_break()
    doc.add_heading('二、模型架构（V2）', level=1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png_arch, width=Inches(6.75))
    caption(doc, '图2　yy25 V2 架构：术数层（矫正轴）→ 体质层（25型）→ 双目标预测引擎 → 耦合/数字人')
    body(doc, '架构自下而上五层：L0 素材 → 术数层（双层进制/60甲子/男8女7/左右镜像，V2 矫正轴）→ '
              '体质层（四类参数 B 与 25 型骨架）→ 预测引擎（目标①健康 / 目标②证候易感 + 辅助引擎）→ '
              '耦合层（Φ_total = ∫Σ[B_体质增益]·R_j(E,P;Z_j)dτ − D0）与数字人接入。')

    # 三、60甲子终身模式
    doc.add_heading('三、60甲子 · 个体终身模式', level=2)
    body(doc, '命局主行 = 本命年天干之五行（如 甲子 → 甲属木 → 命局木）；与先天禀赋五形比对生克，定终身气运基调。')
    make_table(doc, ['命局 vs 禀赋', '基调', '义'],
               [['同', '本命与禀赋同气', '禀性纯粹、易过刚'],
                ['生', '命生禀赋(泄)', '才思外发、易耗散'],
                ['被生', '禀赋生命局(得助)', '根基厚、得扶持'],
                ['克', '命克禀赋', '压力内蕴、易郁'],
                ['被克', '禀赋克命局', '主控力强、易劳心']],
               widths=[30, 45, 85])
    body(doc, '接口：jiazi(index) · jiazi_by_year(year)（甲子参考年 1984）· jiazi_lifetime(index, xing)。', size=9.5,
         color=(0x5D, 0x6D, 0x7E))

    # 四、男8女7
    doc.add_heading('四、男8女7 · 发育阶段性模式', level=2)
    rows = []
    for sex, cn in [('female', '女'), ('male', '男')]:
        st = F.DEV_STAGES[sex]
        rows.append([f'{cn}（{F.LIFE_RHYTHM[sex]["base"]}）', '/'.join(str(x['age']) for x in st),
                     f'{st[1]["age"]}(天癸至)', f'{F.LIFE_RHYTHM[sex]["tiangui_jie"]}(天癸竭)',
                     F.LIFE_RHYTHM[sex]['axis']])
    make_table(doc, ['性别节律', '节点(岁)', '天癸至', '天癸竭', '显性轴'], rows,
               widths=[24, 60, 26, 26, 60])
    body(doc, '阶段序列：童 → 天癸至 → 壮盛 → 巅峰 → 始衰 → 阳衰 → 天癸竭 →（衰极）。接口：development(sex, age)。',
         size=9.5, color=(0x5D, 0x6D, 0x7E))

    # 五、左右镜像
    doc.add_page_break()
    doc.add_heading('五、左右镜像（男女侧别阴阳互易）', level=1)
    body(doc, '矫正规则：男——左为阳、右为阴；女——右为阳、左为阴。', bold=True)
    make_table(doc, ['性别', '左', '右', '阳侧', '脉诊主'],
               [['男', '阳', '阴', '左', '左寸候心/左关候肝/左尺候肾'],
                ['女', '阴', '阳', '右', '右寸候肺/右关候脾/右尺候命门']],
               widths=[16, 12, 12, 14, 106])
    body(doc, '用途：25型亚型方位（如"太角·左足少阳之上"）的阴阳归属随性别翻转；经络取穴、脉诊、发育侧别均按镜像校正。'
              '接口：mirror(sex) · subtype_side_sex(subtype, sex)。', size=9.5)

    # 六、目标1
    doc.add_page_break()
    doc.add_heading('六、预测目标① 个体健康状态', level=1)
    body(doc, 'health_state(xing, sex, jiazi_index, age)：综合 禀赋(五形) + 命局(甲子) + 发育期(男8女7) + 形色相得 → 0–1 健康评分与分级。')
    make_table(doc, ['评分', '状态', '释义'],
               [['≥0.80', '平人', '形色相得、阴阳和平，气血调畅'],
                ['≥0.65', '稳健', '禀赋协调，偶有偏性，宜顺时调摄'],
                ['≥0.50', '亚健康', '禀赋与岁运略失协调，处失衡边缘'],
                ['≥0.35', '失衡', '形色/命局相克或逢衰期，功能失调倾向'],
                ['<0.35', '病态倾向', '多重失和叠加，易发证候，宜早干预']],
               widths=[22, 30, 108])
    ex = m.health_state('火', 'female', 2, 28)
    body(doc, f'算例（火形·女·丙寅命·28岁）：健康评分 {ex["health_score"]} → {ex["state"]}（{ex["meaning"]}）；'
              f'形色{ex["form_color"]}、命局{ex["ming_wuxing"]}、发育期{ex["dev_stage"]}、{ex["mirror"]}。', size=9.5)

    # 七、目标2
    doc.add_heading('七、预测目标② 疾病证候易感性', level=1)
    body(doc, 'syndrome_susceptibility(xing, sex, jiazi_index, age)：五形主证候为基线，叠加 命局五行气旺 + 发育期天癸衰减 + 性别(女以血为用) → 证候易感度排序。')
    make_table(doc, ['五形', '主证候', '病性', '常用兼证'],
               [['木(肝)', '肝郁气滞证', '气滞/血瘀', '肝阳上亢证 · 肝血虚证'],
                ['火(心)', '心火亢盛证', '热/瘀', '心血瘀阻证 · 心阴虚证'],
                ['土(脾)', '脾虚湿困证', '湿/虚', '痰湿内蕴证 · 脾气虚证'],
                ['金(肺)', '肺气虚证', '虚/燥/气郁', '肺燥津伤证 · 肺气郁痹证'],
                ['水(肾)', '肾阳虚证', '虚/寒/水停', '肾精不足证 · 肾水泛滥证']],
               widths=[20, 34, 26, 80])
    sy = m.syndrome_susceptibility('火', 'female', 2, 28)
    body(doc, f'算例（火形·女·丙寅命·28岁）：主证 {sy["main_zheng"]}（病性 {sy["bing_xing"]}）；易感排序 '
              + '；'.join(f'{r["zheng"]} {r["susceptibility"]}({r["level"]})' for r in sy['ranked']) + '。', size=9.5)

    # 八、快速开始 / 文件清单
    doc.add_page_break()
    doc.add_heading('八、快速开始与文件清单', level=1)
    body(doc, '用法（Python）：')
    make_table(doc, ['调用', '功能'],
               [['m.dual_radix()', '阴阳五行双层进制'],
                ['m.jiazi_by_year(1984) / jiazi_lifetime(0,"木")', '60甲子 / 终身模式'],
                ['m.development("female",35)', '男8女7 发育阶段'],
                ['m.mirror("female") / subtype_side_sex(...)', '左右镜像'],
                ['m.health_state("火","female",2,28)', '目标① 健康状态'],
                ['m.syndrome_susceptibility("火","female",2,28)', '目标② 证候易感']],
               widths=[80, 80])
    make_table(doc, ['文件', '作用'],
               [['constitution/framework.py', '符号体系 + 术数层（天干地支/60甲子/生克/男8女7/左右镜像/证候基线/健康分级）'],
                ['constitution/data/wu_xing_25.json', '25型 + 男女参数 + 九类 Hermes 思考 + 疾病易感'],
                ['constitution/model.py', 'ConstitutionModel：生成/逆推/易感/术数/双目标预测'],
                ['constitution/tests/test_constitution.py', 'pytest 套件'],
                ['scripts/gen_yy25v2_docx.py', '本 README → yy25V2.docx 生成器']],
               widths=[64, 96])

    # 九、局限
    doc.add_heading('九、局限与说明', level=1)
    for t in ['术数层为结构化建模：天干地支阴阳五行、60甲子、女七男八、左右镜像均按经典成说编码，不涉及命理吉凶推断；命局"气旺"仅作曲证候加权的启发式信号。',
              '证候易感度为相对排序（0–1 启发值），非临床概率；年忌(7/16/25/34/43/52/61)为高敏感窗提示。',
              '与病因层（etiology/）的运行时耦合、sage-api 接入尚未实现。',
              '25型亚型"质判/判徵"两篇名称略异，并存标注。']:
        p = doc.add_paragraph(style='List Bullet') if 'List Bullet' in [st.name for st in doc.styles] else doc.add_paragraph()
        r = p.add_run(t); r.font.size = Pt(9.5)
        r.font.name = '微软雅黑'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    doc.save(out)
    return out


if __name__ == '__main__':
    pa = os.path.join(TMP, 'yy25v2_arch.png')
    pr = os.path.join(TMP, 'yy25v2_radix.png')
    draw_arch(pa); print('架构图:', pa, os.path.getsize(pa))
    draw_radix(pr); print('进制图:', pr, os.path.getsize(pr))
    out = os.path.join(DESKTOP, 'yy25V2.docx')
    build(pa, pr, out)
    print('docx:', out, os.path.getsize(out))
