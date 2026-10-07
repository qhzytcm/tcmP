# -*- coding: utf-8 -*-
"""中医知识树素材工作簿 · 交付验收（Desktop\\中医知识树.xlsx）。
运行：python scripts/verify_knowledge_tree_xlsx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path.home() / "Desktop" / "中医知识树.xlsx"
fails = []


def check(name, cond, extra=""):
    print(("  ✓ " if cond else "  ✗ ") + name + (f"  [{extra}]" if extra else ""))
    if not cond:
        fails.append(name)


def main():
    print("═" * 62)
    print(" 中医知识树.xlsx · 交付验收")
    print("═" * 62)
    check("文件存在", XLSX.exists(), str(XLSX))
    if not XLSX.exists():
        sys.exit(1)
    wb = load_workbook(XLSX)
    names = wb.sheetnames
    print("  sheets:", len(names))

    need = ["0封面", "0目录", "0正文总表", "第一章 绪论——从知识图谱到中医知识树",
            "第八章 构建管线与工程实现", "第十章 实验", "D07-S04 中医知识树",
            "知识树总览", "图谱接口", "embedding-code", "构建与部署"]
    for n in need:
        check(f"含 sheet「{n}」", n in names)

    # 正文总表
    ws = wb["0正文总表"]
    rows = ws.max_row - 1
    check("0正文总表 含正文行 > 150", rows > 150, rows)
    types = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        types.add(r[2])
    check("正文总表含 标题/正文/代码/表格 四类", {"标题", "正文", "代码", "表格"} <= types, sorted(types))

    # 表格原子行完整性（L4 行必须成整行，验证「超长多行单元格」问题已消除）
    rows_txt = [str(r[3]) for r in ws.iter_rows(min_row=2, values_only=True) if r[3]]
    check("表格原子行完整（L4 行成整行）",
          any(c.strip() == "| L4 | 叶 | 终末知识点／病证单元 | 章节知识点 + 病证单元 |"
              for c in rows_txt))
    check("正文总表无超长多行单元格", all(("\n" not in c) for c in rows_txt))

    # 章节 sheet 有正文
    ws = wb["第八章 构建管线与工程实现"]
    body = "\n".join(str(r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[1])
    check("第八章含 FTS5 代码", "FTS5" in body or "fts5" in body)
    check("第八章含 RRF", "rrf" in body.lower())
    check("第八章含 context engineering", "context" in body.lower())

    # embedding-code
    ws = wb["embedding-code"]
    ec = "\n".join(str(c) for r in ws.iter_rows(values_only=True) for c in r if c)
    for kw in ["ICD", "FTS5", "RRF", "TfidfVectorizer", "tcm_embed.py", "/semantic-search"]:
        check(f"embedding-code 含「{kw}」", kw in ec)

    # 目录结构
    ws = wb["0目录"]
    toc = "\n".join(str(c) for r in ws.iter_rows(values_only=True) for c in r if c)
    check("目录结构含 4 篇 10 章", toc.count("篇") >= 4 and toc.count("章") >= 10)

    # 图谱接口
    ws = wb["图谱接口"]
    gi = "\n".join(str(c) for r in ws.iter_rows(values_only=True) for c in r if c)
    check("图谱接口含 icd11_map", "icd11_map" in gi)
    check("图谱接口含 DSU-00001", "DSU-00001" in gi)

    print("═" * 62)
    if fails:
        print(f" 验收失败 {len(fails)} 项： " + "、".join(fails))
        sys.exit(1)
    print(" 验收全部通过 ✓")


if __name__ == "__main__":
    main()
