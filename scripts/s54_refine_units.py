# -*- coding: utf-8 -*-
"""
S54 精修②：段标记归一 + 合并单元再切分 → 三段齐备率提升
① 段标记变体归一（病达/病然/…→病状；浩法/洛法/…→治法）
② 对「含多个病源且其前已有治法」的单元再切分（段长下限 20，避免碎片）
③ 输出 data/pdf3_units_v2.json + 重出 03-bycd-三段结构.docx
"""
import os, re, json, sys
from collections import Counter
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ocr_fix_rules import fix

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, 'dist')
DESKTOP = r'C:\Users\DELL\Desktop'
ORD = ('病源', '病状', '治法')
RX_MARK = re.compile('病源|病状|治法')
MINLEN = 20


def units_of(S):
    # 归一被 OCR 误认的词目标记：『（名】』 → 『【名】』（左括号误认）
    S = re.sub(r'（([\u4e00-\u9fff]{1,8})】', r'【\1】', S)
    idx = [m.start() for m in re.finditer('【', S)]
    out = []
    for i, p in enumerate(idx):
        end = idx[i + 1] if i + 1 < len(idx) else len(S)
        seg = S[p + 1:end]
        j = seg.find('】')
        if 0 <= j <= 10:
            nm, b = seg[:j], seg[j + 1:]
        else:
            m = re.match(r'([\u4e00-\u9fff]{1,8}?)(?=[，,。；;：:（(0-9０-９]|病源|病状|治法)', seg)
            nm, b = (m.group(1), seg[m.end():]) if m else (seg[:5], seg[5:])
        nm = nm.strip()
        if nm and len(nm) <= 12:
            out.append([nm, b])
    return out


def resplit(us, minlen=MINLEN):
    res = []
    for nm, b in us:
        pos = [m.start() for m in re.finditer('病源', b)]
        cuts = [p for p in pos[1:] if '治法' in b[:p]]
        if not cuts:
            res.append([nm, b]); continue
        segs = []; prev = 0
        for c in cuts:
            segs.append(b[prev:c]); prev = c
        segs.append(b[prev:])
        if all(len(s) >= minlen for s in segs):
            for k, s in enumerate(segs):
                res.append([nm if k == 0 else '病名待定', s])
        else:
            res.append([nm, b])
    return res


def sections(b):
    pos = {}
    for k in ORD:
        m = re.search(k, b)
        if m:
            pos[k] = m.start()
    if not pos:
        return {'全': b}
    order = sorted(pos.items(), key=lambda kv: kv[1])
    out = {}
    for i, (k, p) in enumerate(order):
        e = order[i + 1][1] if i + 1 < len(order) else len(b)
        s = b[p:e]
        if s.startswith(k):
            s = s[len(k):]
        out[k] = s
    if order[0][1] > 0:
        out[order[0][0]] = b[:order[0][1]] + out[order[0][0]]
    return out


def main():
    raw = json.load(open(os.path.join(DIST, 'pdf3_stream.json'), encoding='utf-8'))['full']
    S = fix(raw)                                  # 含 t2s + 校字 + 段标记归一
    json.dump({'full': S}, open(os.path.join(DIST, 'pdf3_stream_v2.json'), 'w', encoding='utf-8'), ensure_ascii=False)

    us0 = units_of(S)
    comp0 = sum(1 for _, b in us0 if all(k in b for k in ORD))
    us = resplit(us0)
    # 依《病名目录》权威序归位（用户逐条校验）：二画段首五名
    HEAD_ORDER = ['丁疤', '丁奚疳', '丁痂皮', '七疝', '七星赶月疔']
    pos = {}
    for i, (nm, b) in enumerate(us):
        if nm in HEAD_ORDER and nm not in pos:
            pos[nm] = i
    if len(pos) >= 2:
        idxs = sorted(pos.values())
        picks = [us[i] for i in idxs]
        picks.sort(key=lambda t: HEAD_ORDER.index(t[0]))
        for k, i in enumerate(idxs):
            us[i] = picks[k]
    units = []
    for nm, b in us:
        sec = sections(b)
        units.append({'name': nm, 'body': b, 'sections': sec,
                      'has': {k: (k in sec) for k in ORD}})
    # S66 词目名定向校正 + 续文合并（用户逐条校验）
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import s66_name_fix
    n_before = len(units)
    units, n_merge, n_fix, remap = s66_name_fix.apply(units)
    json.dump({'remap': {str(k): v for k, v in remap.items()}},
              open(os.path.join(DIST, 'pdf3_units_remap.json'), 'w', encoding='utf-8'))
    print(f'S66: 合并续文 {n_merge} | 改名 {n_fix} | 单元 {n_before} → {len(units)}')
    comp = sum(1 for u in units if all(u['has'].values()))
    st = Counter()
    for u in units:
        for k in ORD:
            if u['has'][k]:
                st[k] += 1
    rep = {'units_before': len(us0), 'complete_before': comp0,
           'rate_before': round(comp0 / max(1, len(us0)) * 100, 1),
           'units': len(units), 'complete': comp,
           'rate': round(comp / max(1, len(units)) * 100, 1),
           'sections': dict(st), 'minlen': MINLEN}
    json.dump(units, open(os.path.join(ROOT, 'data', 'pdf3_units_v2.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=0)
    json.dump(rep, open(os.path.join(DIST, 'pdf3_units_v2_report.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)
    print('段标记归一后：单元', len(us0), '| 齐备', comp0, f'({rep["rate_before"]}%)')
    print('再切分后：  单元', len(units), '| 齐备', comp, f'({rep["rate"]}%)')
    print('分段统计:', dict(st))

    # 重出 docx
    doc = Document()
    s = doc.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297)
    stl = doc.styles['Normal']; stl.font.name = '仿宋'; stl.font.size = Pt(10.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('《病源辞典》病名单元·三段结构重建（bycd/03 清晰本·校字+段标记归一+再切分）')
    r.bold = True; r.font.size = Pt(15)
    q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.add_run(f'共 {len(units)} 个单元　三段齐备 {comp}（{rep["rate"]}%）　'
              f'病源 {st["病源"]} · 病状 {st["病状"]} · 治法 {st["治法"]}').font.size = Pt(10)
    # ── 病名目录（扫描12–58 = 目录01–47，每页 4 栏）──
    toc_fp = os.path.join(DIST, 'pdf3_toc_stream.json')
    if os.path.exists(toc_fp):
        th = doc.add_paragraph(); tr = th.add_run('病名目录（原书 目录01–47，每页 4 栏）')
        tr.bold = True; tr.font.size = Pt(14)
        for it in json.load(open(toc_fp, encoding='utf-8'))['pages']:
            scan = it['scan']
            doc.add_paragraph(f'［目录{scan - 11:02d}·扫描{scan}］ {it["text"]}')
    bh = doc.add_paragraph(); br = bh.add_run('正文（原书 正文0001–1121，每页 3 栏）')
    br.bold = True; br.font.size = Pt(14)
    for i, u in enumerate(units, 1):
        h = doc.add_paragraph(); rr = h.add_run(f'{i:04d}　{u["name"]}')
        rr.bold = True; rr.font.size = Pt(12)
        sec = u['sections']
        if not sec:
            doc.add_paragraph('（无正文内容）'); continue
        for k in ORD:
            if k in sec:
                doc.add_paragraph(f'[{k}]：{sec[k]}')
        for k, v in sec.items():
            if k not in ORD:
                doc.add_paragraph(v)
    out = os.path.join(DESKTOP, '03-bycd-三段结构.docx')
    doc.save(out)
    print('→', out, os.path.getsize(out), 'bytes')


if __name__ == '__main__':
    main()
