# -*- coding: utf-8 -*-
"""cmrl_g4_check.py — G4 出版验收检查（全书完整性门禁）

验收项:
  A. 10 章 sheet 完整: 每章 48 段正文 + 制品区([习题]/[配图]/[自检] 行数>0)
  B. 0cmrl目录: 815 行、四级编码分列(章10/节80/小节240/段落480)
  C. 目录文件: toc_v2/ 10 份 chNN_toc.md + 00-审核目录.md
  D. 制品文件: 每章 chapter-NN-exercises/figures/selfcheck.md + check_chNN.py
  E. 修改说明书: 每章 chapter-NN-changelog.md（修订痕迹）
  F. 三审/复审报告: l3-review/ 下 3 份评审 + 3 份复审 + 汇总
  G. 全书字数汇总
用法: python scripts/cmrl_g4_check.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from cmrl_merge import CHAP_META  # noqa: E402
from cmrl_update_xlsx import XLSX, load_chapter  # noqa: E402

DRAFTS = Path(r'C:\Users\DELL\textbook-project\drafts\cmrl')
REVIEWS = DRAFTS / 'l3-review'


def main():
    from openpyxl import load_workbook
    ok = True
    print('===== G4 出版验收 =====')
    # A. 章节 sheet
    wb = load_workbook(XLSX, read_only=True)
    total_words = 0
    for ch in range(1, 11):
        prefix, sheet_name, title = CHAP_META[ch]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        mods = {}
        for r in rows[2:]:
            m = str(r[1] or '')
            mods[m] = mods.get(m, 0) + 1
        body = sum(v for m, v in mods.items() if m and not m.startswith('['))
        ex, fi, sc = mods.get('[习题]', 0), mods.get('[配图]', 0), mods.get('[自检]', 0)
        # 正文区段数 = 非制品模块行数
        a_ok = body >= 48 and ex > 0 and fi > 0 and sc > 0
        ok &= a_ok
        print(f'  [{"PASS" if a_ok else "FAIL"}] {sheet_name}: 正文段{body} 习题{ex} 配图{fi} 自检{sc}')
        # 字数: 正文区
        w = sum(len(str(r[2] or '')) for r in rows[2:2 + 48])
        total_words += w
    wb.close()
    # B. 目录
    wb = load_workbook(XLSX, read_only=True)
    trows = list(wb['0cmrl目录'].iter_rows(values_only=True))
    wb.close()
    data = trows[5:]
    n_ch = sum(1 for r in data if str(r[0] or '').strip() and not str(r[1] or '').strip()
               and not str(r[2] or '').strip() and not str(r[3] or '').strip())
    n_sec = sum(1 for r in data if str(r[1] or '').strip() and not str(r[2] or '').strip()
                and not str(r[3] or '').strip())
    n_sub = sum(1 for r in data if str(r[2] or '').strip() and not str(r[3] or '').strip())
    n_para = sum(1 for r in data if str(r[3] or '').strip())
    b_ok = len(trows) == 815 and (n_ch, n_sec, n_sub, n_para) == (10, 80, 240, 480)
    ok &= b_ok
    print(f'  [{"PASS" if b_ok else "FAIL"}] 0cmrl目录: {len(trows)}行 章{n_ch}/节{n_sec}/小节{n_sub}/段落{n_para}')
    # C. 目录文件
    c_ok = all((DRAFTS / 'toc_v2' / f'ch{i:02d}_toc.md').exists() for i in range(1, 11)) \
        and (DRAFTS / 'toc_v2' / '00-审核目录.md').exists()
    ok &= c_ok
    print(f'  [{"PASS" if c_ok else "FAIL"}] toc_v2: 10 份 chNN_toc.md + 00-审核目录.md')
    # D. 制品文件
    d_ok = True
    for ch in range(1, 11):
        d = DRAFTS / f'ch{ch:02d}'
        for name in (f'chapter-{ch:02d}-exercises.md', f'chapter-{ch:02d}-figures.md',
                     f'chapter-{ch:02d}-selfcheck.md', f'check_ch{ch:02d}.py'):
            if not (d / name).exists():
                d_ok = False
                print(f'  [MISS] ch{ch:02d}/{name}')
    ok &= d_ok
    print(f'  [{"PASS" if d_ok else "FAIL"}] 制品文件: 10 章 × 4 (3制品+复验脚本)')
    # E. changelog
    e_ok = True
    for ch in range(1, 11):
        if not (DRAFTS / f'ch{ch:02d}' / f'chapter-{ch:02d}-changelog.md').exists():
            e_ok = False
            print(f'  [MISS] ch{ch:02d} changelog')
    ok &= e_ok
    print(f'  [{"PASS" if e_ok else "FAIL"}] 修改说明书: 10 章 changelog')
    # F. 评审报告
    need = ['主编-终审意见.md', '中医药主审-评审意见.md', '智能主审-评审意见.md',
            '主编-复审意见.md', '中医药主审-复审意见.md', '智能主审-复审意见.md',
            '三审汇总.md']
    f_ok = all((REVIEWS / n).exists() for n in need)
    ok &= f_ok
    print(f'  [{"PASS" if f_ok else "FAIL"}] l3-review: 3评审+3复审+汇总')
    # G. 字数
    print(f'  [INFO] 全书正文合计 ≈ {total_words:,} 字')
    print(f'===== G4 验收: {"通过" if ok else "未通过"} =====')
    sys.exit(0 if ok else 1)


if __name__ == '__main__':
    main()
