# -*- coding: utf-8 -*-
"""cmrl_append_assets.py — 把章节配套制品（习题集/配图说明/自检清单）追加到 Excel 章节 sheet

制品 md: drafts/cmrl/ch{NN}/chapter-{NN}-exercises.md / -figures.md / -selfcheck.md
追加: 正文之后, 模块列标 [习题]/[配图]/[自检], 序号续编, 每行一段(表格行/编号题/清单项)

用法: python scripts/cmrl_append_assets.py --ch 1 [--all]
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from cmrl_merge import CHAP_META  # noqa: E402

DRAFTS = Path(r'C:\Users\DELL\textbook-project\drafts\cmrl')
XLSX = Path(r'C:\Users\DELL\Desktop\强化学习的数学原理-赵世钰.xlsx')

ASSETS = {
    'exercises': ('[习题]', 'chapter-{ch:02d}-exercises.md'),
    'figures': ('[配图]', 'chapter-{ch:02d}-figures.md'),
    'selfcheck': ('[自检]', 'chapter-{ch:02d}-selfcheck.md'),
}


def parse_asset(text: str):
    """按行解析制品 md：跳过 # 标题行、--- 分隔行、表格分隔行(|---|)，
    保留表格行/编号题/清单项"""
    out = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith('#') or re.match(r'^-{3,}$', s) \
                or re.match(r'^\|[\s\-|:]+\|$', s):
            continue
        out.append(s)
    return out


def append_assets(ch: int, dry: bool = False):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    prefix, sheet_name, title = CHAP_META[ch]
    d = DRAFTS / f'ch{ch:02d}'
    rows_to_add = []          # (mod, content)
    for key, (mod, fname) in ASSETS.items():
        f = d / fname.format(ch=ch)
        if not f.exists():
            print(f'[warn] 缺 {fname.format(ch=ch)}')
            continue
        lines = parse_asset(f.read_text(encoding='utf-8'))
        rows_to_add.extend((mod, ln) for ln in lines)
    if not rows_to_add:
        print(f'[ch{ch:02d}] 无制品可追加')
        return 0
    wb = load_workbook(XLSX)
    ws = wb[sheet_name]
    # 幂等: 已含制品(存在 [习题]/[配图]/[自检] 模块行)则跳过, 防批量重复追加
    existing = {str(r[1]) for r in ws.iter_rows(min_row=2, max_col=2, values_only=True)}
    if any(m in existing for m in ('[习题]', '[配图]', '[自检]')):
        print(f'[ch{ch:02d}] {sheet_name} 已含制品, 跳过（幂等）')
        wb.close()
        return 0
    start = ws.max_row + 1
    for i, (mod, content) in enumerate(rows_to_add, start):
        ws.append([i, mod, content])
        ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(i, 2).font = Font(bold=True, color='C00000')
        ws.cell(i, 2).fill = PatternFill('solid', fgColor='FFF2CC')
    wb.save(XLSX)
    print(f'[ch{ch:02d}] {sheet_name}: 追加 {len(rows_to_add)} 行制品'
          f'（习题{sum(1 for m, _ in rows_to_add if m == "[习题]")}'
          f'/配图{sum(1 for m, _ in rows_to_add if m == "[配图]")}'
          f'/自检{sum(1 for m, _ in rows_to_add if m == "[自检]")}）→ 行 {start}-{ws.max_row}')
    return len(rows_to_add)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--ch', type=int)
    ap.add_argument('--all', action='store_true')
    args = ap.parse_args()
    chs = range(1, 11) if args.all else [args.ch]
    total = 0
    for ch in chs:
        total += append_assets(ch)
    print(f'合计追加 {total} 行')
