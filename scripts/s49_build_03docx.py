# -*- coding: utf-8 -*-
"""
S49 03 扫描件重建 → 03-bycd.docx
保持原始排版：**逐页、按原书阅读序（竖排右→左、列内上→下）** 输出，不改写、不重排
繁→简：opencc t2s；保护「目标不在常用汉字区」的转换（防豆腐块）
用法： python scripts/s49_build_03docx.py
"""
import os, re, json, glob, argparse
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ocr_fix_rules import fix as safe_t2s          # 复用统一校字规则（含 t2s）

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES = os.path.join(ROOT, 'data', 'bingyuan3_ocr')
DESKTOP = r'C:\Users\DELL\Desktop'


def page_text(pg, xthr=45, nbands=3):
    """竖排：条带自上而下 → 条带内列右→左 → 列内字上→下（曾用整页扁平右→左，错）"""
    fp = os.path.join(RES, f'page_{pg:04d}.json')
    if not os.path.exists(fp):
        return None, 0
    d = json.load(open(fp, encoding='utf-8'))
    L = []
    for x in d['lines']:
        xs = [p[0] for p in x['box']]; ys = [p[1] for p in x['box']]
        if x['text'].strip():
            L.append((sum(xs) / 4, sum(ys) / 4, x['text'].strip()))
    if not L:
        return '', 0
    ymin = min(i[1] for i in L); ymax = max(i[1] for i in L)
    span = max(1.0, ymax - ymin)
    parts = []
    for b in range(nbands):
        lo = ymin + span * b / nbands
        seg = ([i for i in L if lo <= i[1] < ymin + span * (b + 1) / nbands]
               if b < nbands - 1 else [i for i in L if i[1] >= lo])
        if not seg:
            continue
        cols = []
        for cx, cy, t in sorted(seg, key=lambda z: -z[0]):
            for c in cols:
                if abs(c['cx'] - cx) < xthr:
                    c['segs'].append((cy, t)); break
            else:
                cols.append({'cx': cx, 'segs': [(cy, t)]})
        cols.sort(key=lambda c: -c['cx'])
        parts.extend(''.join(t for _, t in sorted(c['segs'])) for c in cols)
    return parts, len(parts)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=os.path.join(DESKTOP, '03-bycd.docx'))
    ap.add_argument('--lo', type=int, default=0)
    ap.add_argument('--hi', type=int, default=99999)
    a = ap.parse_args()

    files = sorted(glob.glob(os.path.join(RES, 'page_*.json')))
    print('已 OCR 页:', len(files))
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    st = doc.styles['Normal']
    st.font.name = '仿宋'
    st.font.size = Pt(10.5)
    hp = doc.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run('吴克潜主编《病源辞典》（繁体原本 OCR·繁转简体校文本）')
    r.font.size = Pt(16); r.bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('底本：桌面 bycd/03-病源辞典.pdf（1182 页）　逐页保留原文阅读序　尚未排版').font.size = Pt(10)
    n = 0
    for fp in files:
        pg = int(re.search(r'(\d+)', os.path.basename(fp)).group(1))
        if not (a.lo <= pg <= a.hi):
            continue
        txt, ncol = page_text(pg)
        if not txt:
            continue
        h = doc.add_paragraph()
        rr = h.add_run(f'【第 {pg} 页】（{ncol} 列）')
        rr.bold = True; rr.font.size = Pt(11)
        for col in txt:
            doc.add_paragraph(safe_t2s(col))
        n += 1
    doc.save(a.out)
    print('写入页数:', n, '→', a.out, os.path.getsize(a.out), 'bytes')


if __name__ == '__main__':
    main()
