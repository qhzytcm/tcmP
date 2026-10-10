# -*- coding: utf-8 -*-
"""《中医智能仪器与可穿戴设备》(D07-S11) 教材素材工作簿生成器
GitBook 书稿 → Excel（目录结构 + 正文 + 真实器械数据 + Holter 实测 + embedding-code）。
产出：桌面 中医智能仪器与可穿戴设备.xlsx
"""
from __future__ import annotations
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from gitbook_to_excel import SW, parse_summary, sanitize, write_book  # noqa: E402

REPO = SCRIPTS.parent
BOOK = REPO / "kg" / "device" / "book"
CAT = REPO / "kg" / "device" / "tcm-device.json"
OUT = Path.home() / "Desktop" / "中医智能仪器与可穿戴设备.xlsx"

F_T = Font(name="微软雅黑", size=13, bold=True, color="1F3864")
F_H = Font(name="微软雅黑", size=10.5, bold=True, color="2E5C8A")
F_B = Font(name="微软雅黑", size=10)
F_C = Font(name="Consolas", size=9.5)
AL_L = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_data_sheets(wb, cat, used):
    def new(name, w):
        return SW(wb.create_sheet(sanitize(name, used)), w)

    sw = new("器械分类", [8, 18, 44, 16]); sw.title("六类器械（C1–C6）", 4)
    sw.row(["编号", "类别", "代表器械", "信号"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for c in cat["classes"]:
        sw.row([c["id"], c["name"], c["examples"], c["signal"]], [None] * 4, [F_B] * 4, [AL_C, AL_L, AL_L, AL_C], 18)

    sw = new("可穿戴形态", [8, 16, 30, 30]); sw.title("六种可穿戴形态与中医对应", 4)
    sw.row(["编号", "形态", "传感器", "中医语义"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for f in cat["forms"]:
        sw.row([f["id"], f["name"], "·".join(f["sensors"]), f["tcm"]], [None] * 4, [F_B] * 4, [AL_C, AL_L, AL_L, AL_L], 18)

    sw = new("信号通道", [8, 16, 14, 46]); sw.title("六条生理信号通道", 4)
    sw.row(["编号", "通道", "采样率", "可提取特征"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for s in cat["signals"]:
        sw.row([s["id"], s["name"], s["fs"], s["feat"]], [None] * 4, [F_B] * 4, [AL_C, AL_C, AL_C, AL_L], 18)

    sw = new("数据标准", [22, 42, 12]); sw.title("数据标准与平台落地", 3)
    sw.row(["标准", "作用", "状态"], [None] * 3, [F_H] * 3, [AL_C] * 3, 20)
    for s in cat["standards"]:
        sw.row([s["code"], s["name"], s["verify"]], [None] * 3, [F_B] * 3, [AL_C, AL_L, AL_C], 18)

    sw = new("AI算子", [8, 16, 34, 40]); sw.title("六类器材信息学算子", 4)
    sw.row(["编号", "算子", "算法", "用途"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for o in cat["operators"]:
        sw.row([o["id"], o["name"], o["algo"], o["use"]], [None] * 4, [F_B] * 4, [AL_C, AL_L, AL_L, AL_L], 18)

    sw = new("平台接入", [34, 40]); sw.title("械者 六端点（sage-api 已上线）", 2)
    sw.row(["接口", "用途"], [None] * 2, [F_H] * 2, [AL_C] * 2, 20)
    for p in cat["platform"]:
        sw.row([p["api"], p["use"]], [None] * 2, [F_B] * 2, [AL_L, AL_L], 18)

    sw = new("Holter实测", [14, 10, 10, 10, 10, 10, 34]); sw.title("Holter 心律检测 · 六场景实测（仿真 RR 序列）", 7)
    sw.row(["场景", "HR", "SDNN", "RMSSD", "pNN50", "CV", "判读 + 中医提示"], [None] * 7, [F_H] * 7, [AL_C] * 7, 20)
    for k, v in cat["holter"].items():
        m = v["metrics"]
        sw.row([k, m["hr"], m["sdnn"], m["rmssd"], m["pnn50"], m["cv"],
                ",".join(v["diagnoses"]) + " ｜ " + "；".join(v["tcm_hints"])],
               [None] * 7, [F_B] * 7, [AL_C, AL_C, AL_C, AL_C, AL_C, AL_C, AL_L], 18)


EMBED = [
 ("① 时域 HRV（kg/device/holter.py）",
  "# RR 间期序列(ms) → 指标\nrr = [float(x) for x in rr_ms if x and x > 0]\nmean = sum(rr)/len(rr);  hr = 60000.0/mean\nsdnn  = sqrt(sum((x-mean)**2 for x in rr)/(len(rr)-1))\nrmssd = sqrt(sum(d*d for d in diffs)/len(diffs))\npnn50 = 100.0*sum(1 for d in diffs if abs(d) > 50)/len(diffs)"),
 ("② 节律判读与危险度排序",
  "if hr > 100: flags.append(('SINUS_TACHYCARDIA','warning'))\nif any(x >= 2500 for x in rr): flags.append(('PAUSE','danger'))\nif rmssd > 120 and pnn50 > 45 and cv > 0.12:\n    flags.append(('AF_SUSPECT','warning'))\nflags.sort(key=lambda f: order[f[1]])   # danger 先出"),
 ("③ ICD-11 桥接（scripts/icd11_client.py）",
  "import sqlite3\ncon = sqlite3.connect('tcmP/data/icd11_mms.db')   # 31,838 实体\nrow = con.execute('SELECT id, code, title FROM entities WHERE id=?', (fid,)).fetchone()"),
 ("④ 检索 FTS5 + RRF（scripts/tcm_embed.py）",
  "CREATE VIRTUAL TABLE dsu_fts USING fts5(id UNINDEXED, text, tokenize='unicode61');\nSELECT id, bm25(dsu_fts) FROM dsu_fts WHERE dsu_fts MATCH '心悸' ORDER BY 2 LIMIT 5;\n# RRF：score = Σ 1/(k + rank), k=60"),
 ("⑤ 上下文工程装配（context engineering）",
  "def build_context(q, k=5):\n    hits = retrieve(q, k)          # 向量 ∪ FTS5 ∪ RRF\n    return '\\n'.join(f'[{h.id}] {h.disease} / {h.syndrome}（证据 {h.evidence}）' for h in hits)\n# Top-5 ≈ 1.5K token，较全库 ↓≈95%"),
 ("⑥ 中医药关联提示（仅提示，非诊断）",
  "TCM_RULES = [('hr>100 且 rmssd>60', '心火亢盛 / 阴虚火旺'),\n             ('sdnn<30',           '心气虚 / 心脉瘀阻'),\n             ('AF_SUSPECT',        '心悸、怔忡（脉结代/促/涩）')]\n# 输出必须携带边界声明：不构成医学诊断，不得据以调整用药"),
]


def add_embed_sheet(wb, used):
    sw = SW(wb.create_sheet(sanitize("embedding-code", used)), [30, 96])
    sw.title("embedding code：HRV · 节律判读 · ICD-11 · FTS5 · RRF · 上下文工程", 2)
    for t, code in EMBED:
        sw.row([t, ""], [None, None], [F_H, F_B], [AL_L, AL_L], 18)
        for ln in code.split("\n"):
            sw.row(["", ln], [None, None], [F_B, F_C], [AL_L, AL_L], 15 if ln.strip() else 12)


def main():
    cat = json.loads(CAT.read_text(encoding="utf-8"))
    wb = Workbook(); wb.remove(wb.active); used = set()
    meta = dict(title="中医智能仪器与可穿戴设备", subtitle="TCM Smart Instruments & Wearables",
                author="六者·中医医院AI网络教育平台 · tcmP",
                note="D07-S11 教材素材 · 与 kg/device 单一构建点同步")
    write_book(wb, BOOK, meta, used)
    add_data_sheets(wb, cat, used)
    add_embed_sheet(wb, used)
    wb.save(OUT)
    print("[产出] " + str(OUT))
    print("[sheet] %d 个：%s" % (len(wb.sheetnames), ", ".join(wb.sheetnames)))


if __name__ == "__main__":
    main()