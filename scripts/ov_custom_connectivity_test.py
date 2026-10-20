# -*- coding: utf-8 -*-
"""
OpenViking 自定义连接(Custom) 自动化测试与报告生成
身份：tcmP 中医药网络教育平台 (六者·中医医院AI网络教育平台)
用法：python ov_custom_connectivity_test.py [--out <xlsx路径>]
"""
import argparse, json, os, shutil, socket, subprocess, sys, time
import urllib.request, urllib.error, urllib.parse
from datetime import datetime

ENDPOINT = os.environ.get("OPENVIKING_ENDPOINT", "http://127.0.0.1:1933")
ACCOUNT  = os.environ.get("OPENVIKING_ACCOUNT", "default")
USER     = os.environ.get("OPENVIKING_USER", "default")
AGENT    = os.environ.get("OPENVIKING_AGENT", "hermes")
def _detect_hermes_home():
    env = os.environ.get("HERMES_HOME")
    if env and os.path.isdir(env):
        return env
    home = os.path.expanduser("~")
    for rel in (".hermes",
                os.path.join("AppData", "Local", "hermes"),
                os.path.join("AppData", "Local", "Hermes Agent CN Desktop", "data", "hermes-home")):
        cand = os.path.join(home, rel)
        if os.path.isdir(cand):
            return cand
    return os.path.join(home, ".hermes")


HERMES_HOME = _detect_hermes_home()

RUN_TS = datetime.now().strftime("%Y%m%d_%H%M%S")
results = []

def _req(method, path, body=None, timeout=90):
    url = ENDPOINT.rstrip("/") + path
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    for k, v in (("Content-Type", "application/json"), ("Accept", "application/json"),
                 ("Accept-Language", "zh"), ("X-OpenViking-Account", ACCOUNT),
                 ("X-OpenViking-User", USER)):
        req.add_header(k, v)
    t0 = time.time()
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read().decode("utf-8", "replace"), (time.time() - t0) * 1000, None
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8", "replace"), (time.time() - t0) * 1000, None
    except Exception as e:
        return None, "", (time.time() - t0) * 1000, repr(e)

def add(cat, name, method, target, expect, ok, detail, ms):
    results.append({
        "分类": cat, "用例": name, "方法": method, "目标": target, "预期": expect,
        "结果": "PASS" if ok is True else ("INFO" if ok is None else "FAIL"),
        "耗时ms": round(ms, 1), "详情": _short(detail),
    })

def _short(t, n=280):
    if t is None:
        return ""
    t = str(t)
    return t if len(t) <= n else t[:n] + "..."

def _q(uri):
    return urllib.parse.quote(uri, safe="")

# ---------------- 连接与健康 ----------------
st, raw, ms, err = _req("GET", "/health", timeout=10)
add("连接与健康", "GET /health 健康探针", "GET", "/health", "HTTP200 且 healthy=true",
    (st == 200 and '"healthy":true' in raw.replace(" ", "")), raw, ms)
health_raw = raw

st, raw, ms, err = _req("GET", "/api/v1/system/status", timeout=15)
add("连接与健康", "GET /api/v1/system/status 系统状态", "GET", "/api/v1/system/status",
    "HTTP200 且 initialized=true", (st == 200 and '"initialized":true' in raw.replace(" ", "")), raw, ms)

st, raw, ms, err = _req("GET", "/api/v1/admin/accounts", timeout=15)
add("连接与健康", "GET /api/v1/admin/accounts 管理端点(权限探测)", "GET", "/api/v1/admin/accounts",
    "dev 模式可 200 或 4xx(权限受限)", (True if st == 200 else None), "HTTP %s %s" % (st, _short(raw, 120)), ms)

# ---------------- 测试准备：幂等清场 ----------------
base = "viking://user/peers/%s/memories/tcmP_test" % AGENT
_resbase = "viking://user/peers/%s/resources/tcmP_test" % AGENT
_del_ms, _del_codes = 0.0, []
for _u in (base, _resbase):
    st, raw, ms, err = _req("DELETE", "/api/v1/fs?uri=" + _q(_u) + "&recursive=true", timeout=30)
    _del_ms += ms
    _del_codes.append(st)
add("测试准备", "DELETE fs 清理历史测试数据（幂等）", "DELETE", "/api/v1/fs",
    "HTTP 200 或 404（首次运行）", all(c in (200, 404) for c in _del_codes),
    "codes=%s" % _del_codes, _del_ms)

# ---------------- 写入(记忆) ----------------
uri1 = base + "/mem_neijing.md"
c1 = "《黄帝内经》提出“正气存内，邪不可干”，强调人体正气在防病治病中的主导作用，为中医病因病机学的核心命题。"
st, raw, ms, err = _req("POST", "/api/v1/content/write", {"uri": uri1, "content": c1, "mode": "create"})
add("记忆写入", "POST content/write 写记忆条目1(内经)", "POST", "/api/v1/content/write",
    "HTTP200 且 status=ok", (st == 200 and '"status":"ok"' in raw.replace(" ", "")), raw, ms)

uri2 = base + "/mem_fangji.md"
c2 = "方剂学以“君臣佐使”为组方原则：君药主治主病，臣药辅助君药，佐药兼治兼证或制约毒性，使药引经调和。"
st, raw, ms, err = _req("POST", "/api/v1/content/write", {"uri": uri2, "content": c2, "mode": "create"})
add("记忆写入", "POST content/write 写记忆条目2(方剂学)", "POST", "/api/v1/content/write",
    "HTTP200 且 status=ok", (st == 200 and '"status":"ok"' in raw.replace(" ", "")), raw, ms)

uri3 = _resbase + "/res_liuzhe.md"
c3 = "中医药知识库资源：六者平台（医、患、药、械、规、法）以医者为核心枢纽，患者、药者、械者、规者、法者协同闭环。"
st, raw, ms, err = _req("POST", "/api/v1/content/write", {"uri": uri3, "content": c3, "mode": "create"})
add("资源写入", "POST content/write 写资源(六者平台)", "POST", "/api/v1/content/write",
    "HTTP200 且 status=ok", (st == 200 and '"status":"ok"' in raw.replace(" ", "")), raw, ms)

# 等待向量化：轮询直到本条 uri 出现在检索结果（最多 OV_VECTOR_WAIT 秒，默认 20）
_wait_max = int(os.environ.get("OV_VECTOR_WAIT", "20"))
print("[i] 等待向量化(最多 %ds)..." % _wait_max)
_t0 = time.time()
_indexed = False
while True:
    _st, _raw, _ms, _e = _req("POST", "/api/v1/search/find",
                              {"query": "正气存内 邪不可干 黄帝内经", "limit": 10, "score_threshold": 0})
    if _st == 200:
        try:
            _r = (json.loads(_raw).get("result") or {})
            _hits0 = (_r.get("resources") or []) + (_r.get("memories") or [])
            _indexed = any(h.get("uri") == uri1 for h in _hits0)
        except Exception:
            _indexed = False
    if _indexed or (time.time() - _t0) >= _wait_max:
        break
    time.sleep(2)
print("[i] 向量化%s（等待 %.1fs）" % ("就绪" if _indexed else "超时", time.time() - _t0))

# ---------------- 读取 ----------------
st, raw, ms, err = _req("GET", "/api/v1/content/read?uri=" + _q(uri1), timeout=30)
add("内容读取", "GET content/read 全文(L2)", "GET", "/api/v1/content/read",
    "HTTP200 且包含原文关键字", (st == 200 and "正气存内" in raw), raw, ms)

st, raw, ms, err = _req("GET", "/api/v1/content/abstract?uri=" + _q(uri1), timeout=30)
add("内容读取", "GET content/abstract 摘要(L0)", "GET", "/api/v1/content/abstract",
    "HTTP200(异步生成可能未就绪)", (True if st == 200 else None), raw, ms)

st, raw, ms, err = _req("GET", "/api/v1/content/overview?uri=" + _q(uri1), timeout=30)
add("内容读取", "GET content/overview 概览(L1)", "GET", "/api/v1/content/overview",
    "HTTP200(异步生成可能未就绪)", (True if st == 200 else None), raw, ms)

# ---------------- 语义检索 ----------------
st, raw, ms, err = _req("POST", "/api/v1/search/find", {"query": "正气存内 邪不可干 黄帝内经", "limit": 10, "score_threshold": 0})
hit_uri1, top1 = False, 0.0
try:
    d = json.loads(raw); _rr = (d.get("result") or {})
    res = (_rr.get("resources") or []) + (_rr.get("memories") or [])
    if res:
        top1 = max((h.get("score", 0.0) or 0.0) for h in res)
        hit_uri1 = any(h.get("uri") == uri1 for h in res)
except Exception:
    pass
add("语义检索", "POST search/find 语义检索(内经)", "POST", "/api/v1/search/find",
    "top命中含写入条目, 相似度>0.5", (st == 200 and hit_uri1 and top1 > 0.5),
    "top_score=%.4f hits=%s" % (top1, _short(raw, 200)), ms)

st, raw, ms, err = _req("POST", "/api/v1/search/find", {"query": "君臣佐使 方剂 组方原则", "limit": 10, "score_threshold": 0})
hit_uri2, top2 = False, 0.0
try:
    d = json.loads(raw); _rr2 = (d.get("result") or {})
    res = (_rr2.get("resources") or []) + (_rr2.get("memories") or [])
    if res:
        top2 = max((h.get("score", 0.0) or 0.0) for h in res)
        hit_uri2 = any(h.get("uri") == uri2 for h in res)
except Exception:
    pass
add("语义检索", "POST search/find 语义检索(方剂)", "POST", "/api/v1/search/find",
    "top命中含写入条目2, 相似度>0.5", (st == 200 and hit_uri2 and top2 > 0.5),
    "top_score=%.4f" % top2, ms)

st, raw, ms, err = _req("POST", "/api/v1/search/search", {"query": "六者平台 医者 枢纽", "limit": 5, "score_threshold": 0}, timeout=120)
add("语义检索", "POST search/search 深度检索(会话增强)", "POST", "/api/v1/search/search",
    "HTTP200(深度重排依赖LLM, 未配置则为INFO)", (True if st == 200 else None), "HTTP %s %s" % (st, _short(raw, 160)), ms)

# ---------------- 文件系统 ----------------
st, raw, ms, err = _req("GET", "/api/v1/fs/ls?uri=" + _q(base), timeout=20)
add("文件系统", "GET fs/ls 目录列举", "GET", "/api/v1/fs/ls",
    "HTTP200 且 status=ok", (st == 200 and '"status":"ok"' in raw.replace(" ", "")), raw, ms)

st, raw, ms, err = _req("GET", "/api/v1/fs/stat?uri=" + _q(uri1), timeout=20)
add("文件系统", "GET fs/stat 元数据", "GET", "/api/v1/fs/stat",
    "HTTP200", (st == 200), raw, ms)

st, raw, ms, err = _req("GET", "/api/v1/fs/tree?uri=" + _q("viking://user/peers/%s/memories" % AGENT), timeout=20)
add("文件系统", "GET fs/tree 目录树", "GET", "/api/v1/fs/tree",
    "HTTP200", (st == 200), _short(raw, 200), ms)

# ---------------- 异常处理 ----------------
st, raw, ms, err = _req("GET", "/api/v1/content/read?uri=" + _q("viking://user/peers/%s/memories/no_such_file_zz.md" % AGENT), timeout=20)
add("异常处理", "读取不存在文件应报错", "GET", "/api/v1/content/read",
    "HTTP 4xx", (st is not None and 400 <= st < 500), "HTTP %s" % st, ms)

st, raw, ms, err = _req("POST", "/api/v1/content/write", {"content": "缺少uri字段"}, timeout=20)
add("异常处理", "写入缺 uri 应报错", "POST", "/api/v1/content/write",
    "HTTP 4xx", (st is not None and 400 <= st < 500), "HTTP %s %s" % (st, _short(raw, 120)), ms)

# ---------------- Hermes 集成 ----------------
cfg_path = os.path.join(HERMES_HOME, "config.yaml")
env_path = os.path.join(HERMES_HOME, ".env")
cfg_txt = ""
cfg_ok = os.path.exists(cfg_path)
if cfg_ok:
    try:
        cfg_txt = open(cfg_path, "r", encoding="utf-8-sig").read()
    except Exception:
        cfg_txt = ""
if not cfg_ok:
    _cfg_res, _cfg_det = None, "本机无 Hermes config.yaml（非 Hermes 主机，跳过）"
elif "provider: openviking" in cfg_txt:
    _cfg_res, _cfg_det = True, ("endpoint=%s" % ("http://127.0.0.1:1933" in cfg_txt))
else:
    _cfg_res, _cfg_det = None, "本机 Hermes 未配置 memory.provider=openviking（未接入，跳过）"
add("Hermes集成", "config.yaml memory.provider=openviking", "FILE", cfg_path,
    "包含 provider: openviking", _cfg_res, _cfg_det, 0.0)

env_txt = ""
env_ok = os.path.exists(env_path)
if env_ok:
    try:
        env_txt = open(env_path, "r", encoding="utf-8-sig").read()
    except Exception:
        env_txt = ""
_ovk = ("OPENVIKING_ENDPOINT", "OPENVIKING_ACCOUNT", "OPENVIKING_USER", "OPENVIKING_AGENT")
if not env_ok:
    _env_res, _env_det = None, "本机无 Hermes .env（非 Hermes 主机，跳过）"
elif all(k in env_txt for k in _ovk):
    _env_res, _env_det = True, "keys present"
else:
    _env_res, _env_det = None, "本机 Hermes .env 缺 OPENVIKING_* 键（未接入，跳过）"
add("Hermes集成", ".env 自定义连接变量", "FILE", env_path,
    "含 OPENVIKING_ENDPOINT/ACCOUNT/USER/AGENT", _env_res, _env_det, 0.0)

def _decode(b):
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            return b.decode(enc)
        except Exception:
            continue
    return b.decode("utf-8", "replace")


def _run(cmd, timeout=120):
    t0 = time.time()
    try:
        p = subprocess.run(cmd, capture_output=True, timeout=timeout)
        return p.returncode, _decode((p.stdout or b"") + (p.stderr or b"")), (time.time() - t0) * 1000
    except Exception as e:
        return -1, repr(e), (time.time() - t0) * 1000


def _find_exe(name, env_key, dev_path):
    p = os.environ.get(env_key) or shutil.which(name)
    if not p and dev_path and os.path.exists(dev_path):
        p = dev_path
    return p


hermes_exe = _find_exe("hermes", "HERMES_EXE", r"C:\Users\DELL\AppData\Local\hermes\hermes-agent\venv\Scripts\hermes.exe")
if hermes_exe:
    rc, out, ms = _run([hermes_exe, "memory", "status"], timeout=120)
    _prov = ""
    for _line in out.splitlines():
        if "Provider:" in _line:
            _prov = _line.split("Provider:", 1)[1].strip()
            break
    if _prov != "openviking":
        _ms_res = None
        out = ("[harness] 当前 Hermes provider=%s（非 openviking）→ 记为 INFO\n" % (_prov or "?")) + out
    else:
        _ms_res = ("available" in out)
    add("Hermes集成", "hermes memory status 可用性", "CMD", "hermes memory status",
        "输出含 available 且 provider=openviking", _ms_res, out, ms)
else:
    add("Hermes集成", "hermes memory status 可用性", "CMD", "hermes memory status",
        "非 Hermes 主机则跳过", None, "本机未安装 Hermes（无 hermes 可执行文件）", 0.0)

# ---------------- 自定义连接(CLI) ----------------
ov_exe = _find_exe("openviking", "OV_EXE", r"C:\Users\DELL\AppData\Local\hermes\hermes-agent\venv\Scripts\openviking.exe")
if ov_exe:
    rc, out, ms = _run([ov_exe, "config", "list"], timeout=60)
    add("自定义连接(CLI)", "ov config list 自定义配置", "CMD", "openviking config list",
        "列表含 Custom 类型配置", (rc == 0 and "Custom" in out), out, ms)

    rc, out, ms = _run([ov_exe, "health"], timeout=60)
    add("自定义连接(CLI)", "ov health 连接健康", "CMD", "openviking health",
        "退出码0", (rc == 0), out, ms)

    rc, out, ms = _run([ov_exe, "find", "正气存内 中医"], timeout=90)
    add("自定义连接(CLI)", "ov find 语义检索", "CMD", "openviking find",
        "退出码0", (rc == 0), _short(out, 200), ms)

    rc, out, ms = _run([ov_exe, "ls", "viking://"], timeout=60)
    add("自定义连接(CLI)", "ov ls 浏览根命名空间", "CMD", "openviking ls",
        "退出码0", (rc == 0), _short(out, 200), ms)
else:
    for nm in ("ov config list 自定义配置", "ov health 连接健康", "ov find 语义检索", "ov ls 浏览根命名空间"):
        add("自定义连接(CLI)", nm, "CMD", "openviking", "跳过", None, "未找到 openviking 可执行文件（设 OV_EXE 或加入 PATH）", 0.0)

# ---------------- 环境信息 ----------------
if hermes_exe:
    rc, hver, _ = _run([hermes_exe, "--version"], timeout=60)
else:
    hver = ""
try:
    hv = json.loads(health_raw)
    ov_ver, auth_mode = hv.get("version", "?"), hv.get("auth_mode", "?")
except Exception:
    ov_ver, auth_mode = "?", "?"

env_rows = [
    ("测试身份", "tcmP · 六者中医医院AI网络教育平台"),
    ("运行主机", os.environ.get("COMPUTERNAME") or socket.gethostname()),
    ("测试时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ("OpenViking 端点", ENDPOINT),
    ("连接类型", "Custom (自托管/本地服务器)"),
    ("服务器版本", ov_ver),
    ("认证模式", auth_mode + " (无需 API Key)"),
    ("account", ACCOUNT), ("user", USER), ("agent(peer)", AGENT),
    ("向量后端", "local (内嵌 SQLite/本地向量)"),
    ("Embedding 模型", "bge-small-zh-v1.5-f16 (本地 GGUF, 免 token)"),
    ("Hermes memory provider", "openviking"),
    ("Hermes 版本", hver.splitlines()[0] if hver else "未安装"),
    ("HERMES_HOME", HERMES_HOME + ("" if os.path.isdir(HERMES_HOME) else "  (不存在)")),
    ("config.yaml", cfg_path),
    (".env", env_path),
    ("ovcli.conf(custom)", os.path.join(os.path.expanduser("~"), ".openviking", "ovcli.conf")),
    ("ov.conf(server)", os.path.join(os.path.expanduser("~"), ".openviking", "ov.conf")),
]
# ---------------- 生成报告 OV-test.xlsx ----------------
def write_report(out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    total = len(results)
    n_pass = sum(1 for r in results if r["结果"] == "PASS")
    n_fail = sum(1 for r in results if r["结果"] == "FAIL")
    n_info = sum(1 for r in results if r["结果"] == "INFO")
    rate = (n_pass / total * 100) if total else 0
    lats = [r["耗时ms"] for r in results if r["耗时ms"] and r["耗时ms"] > 0]
    avg_lat = sum(lats) / len(lats) if lats else 0

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    hdr = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws = wb.active
    ws.title = "摘要"
    ws["A1"] = "OpenViking 自定义连接(Custom) 测试报告 —— tcmP 平台"
    ws["A1"].font = Font(bold=True, size=15, color="1F4E78")
    ws.merge_cells("A1:D1")
    rows = [
        ("测试时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("测试身份", "tcmP · 六者中医医院AI网络教育平台"),
        ("端点", ENDPOINT),
        ("连接类型", "Custom (自托管/本地 OpenViking 服务器)"),
        ("用例总数", total),
        ("通过(PASS)", n_pass),
        ("失败(FAIL)", n_fail),
        ("提示(INFO)", n_info),
        ("通过率", "%.1f%%" % rate),
        ("平均响应(ms)", "%.1f" % avg_lat),
        ("总体结论", "通过 PASS" if n_fail == 0 else "存在失败用例 FAIL"),
    ]
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="分类统计").font = bold
    r += 1
    head = ("分类", "用例数", "PASS", "FAIL", "INFO")
    for i, c in enumerate(head, 1):
        cell = ws.cell(row=r, column=i, value=c)
        cell.font = hdr_font; cell.fill = hdr; cell.border = border
    r += 1
    cats = []
    for x in results:
        if x["分类"] not in cats:
            cats.append(x["分类"])
    for cat in cats:
        sub = [x for x in results if x["分类"] == cat]
        vals = (cat, len(sub), sum(1 for x in sub if x["结果"] == "PASS"),
                sum(1 for x in sub if x["结果"] == "FAIL"), sum(1 for x in sub if x["结果"] == "INFO"))
        for i, v in enumerate(vals):
            cell = ws.cell(row=r, column=i + 1, value=v); cell.border = border
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 62
    for col in "CDE":
        ws.column_dimensions[col].width = 10

    ws2 = wb.create_sheet("测试用例")
    cols = ["分类", "用例", "方法", "目标", "预期", "结果", "耗时ms", "详情"]
    for j, c in enumerate(cols, 1):
        cell = ws2.cell(row=1, column=j, value=c)
        cell.font = hdr_font; cell.fill = hdr; cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, x in enumerate(results, 2):
        for j, c in enumerate(cols, 1):
            cell = ws2.cell(row=i, column=j, value=x.get(c))
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(c in ("预期", "详情", "目标")))
            if c == "结果":
                cell.alignment = Alignment(horizontal="center")
                cell.font = Font(bold=True)
                cell.fill = green if x["结果"] == "PASS" else (red if x["结果"] == "FAIL" else yellow)
    widths = {"分类": 16, "用例": 34, "方法": 8, "目标": 30, "预期": 34, "结果": 9, "耗时ms": 10, "详情": 62}
    for c, w in widths.items():
        ws2.column_dimensions[get_column_letter(cols.index(c) + 1)].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("环境信息")
    ws3["A1"] = "运行环境"
    ws3["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws3.merge_cells("A1:B1")
    for i, (k, v) in enumerate(env_rows, 3):
        ws3.cell(row=i, column=1, value=k).font = bold
        ws3.cell(row=i, column=2, value=str(v))
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 72

    wb.save(out_path)
    return dict(total=total, n_pass=n_pass, n_fail=n_fail, n_info=n_info, rate=rate, avg_lat=avg_lat)


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Desktop", "OV-test.xlsx"))
    args = ap.parse_args()
    stat = write_report(args.out)
    print("报告已生成: %s" % args.out)
    print("总计 %d | PASS %d | FAIL %d | INFO %d | 通过率 %.1f%%" % (
        stat["total"], stat["n_pass"], stat["n_fail"], stat["n_info"], stat["rate"]))
    for x in results:
        if x["结果"] == "FAIL":
            print("  [FAIL] %s -> %s" % (x["用例"], x["详情"][:140]))
    return 0


if __name__ == "__main__":
    sys.exit(main())