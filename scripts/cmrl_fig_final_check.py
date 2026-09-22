# -*- coding: utf-8 -*-
"""cmrl_fig_final_check.py — 全书图表终检（聚焦重绘版）

检查:
  A. 10 章 PNG 全部为新聚焦版（存在非空）
  B. Excel xl/media 45 PNG 且与磁盘最新 PNG 哈希一致（图片字节已替换）
  C. 每章图题行(居中黑体)/图注行(左顶格四要素) 格式
  D. 图片锚点 = 图注行-1（三行结构对齐）
"""
import hashlib
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from cmrl_merge import CHAP_META  # noqa: E402
from cmrl_update_xlsx import DRAFTS, XLSX  # noqa: E402

FIGDIR = DRAFTS / 'figures'
results = []


def check(name, ok, detail=''):
    results.append((name, ok, detail))
    print(f'{"PASS" if ok else "FAIL"}  {name}  {detail}')


# A. 磁盘 PNG
png_ok = True
total_png = 0
for ch in range(1, 11):
    d = FIGDIR / f'ch{ch:02d}'
    pngs = sorted(d.glob('fig*.png')) if d.exists() else []
    total_png += len(pngs)
    if len(pngs) < 4 or any(p.stat().st_size < 10 * 1024 for p in pngs):
        png_ok = False
        print(f'  [MISS] ch{ch:02d} PNG: {len(pngs)}')
check(f'A. 磁盘新 PNG: 共 {total_png} 幅非空', png_ok, '')

# B. Excel media 字节 = 磁盘最新 PNG（哈希全比对）
X = Path(XLSX)
with zipfile.ZipFile(X) as z:
    media_pngs = [n for n in z.namelist() if n.startswith('xl/media/') and n.endswith('.png')]
    media_hashes = {hashlib.md5(z.read(n)).hexdigest() for n in media_pngs}
disk_hashes = set()
for ch in range(1, 11):
    for f in (FIGDIR / f'ch{ch:02d}').glob('fig*.png'):
        disk_hashes.add(hashlib.md5(f.read_bytes()).hexdigest())
b_ok = len(media_pngs) >= 45 and disk_hashes <= media_hashes
check(f'B. Excel media 45 PNG = 磁盘最新图（哈希全等 {len(disk_hashes & media_hashes)}/45）',
      b_ok, f'media={len(media_pngs)}')

# C/D. 格式 + 锚点
from openpyxl import load_workbook
wb = load_workbook(X)
bad = []
n_img = 0
for ch in range(1, 11):
    prefix, sheet, title = CHAP_META[ch]
    ws = wb[sheet]
    for img in ws._images:
        n_img += 1
        r = img.anchor._from.row + 1
        cap = ws.cell(r - 1, 3)
        note = ws.cell(r + 1, 3)
        cap_txt = str(cap.value or '')
        note_txt = str(note.value or '')
        if not (cap_txt.startswith('图') and cap.font.name == '黑体'
                and cap.alignment.horizontal == 'center'):
            bad.append(f'{sheet}R{r-1}图题')
        # 图注验收: 图注：开头 + 左顶格 + 内容充实(≥100字) + 对接正文提示语
        note_ok = (note_txt.startswith('图注：')
                   and note.alignment.horizontal == 'left'
                   and len(note_txt) >= 100
                   and any(k in note_txt for k in ('正文', '对接', '配合')))
        if not note_ok:
            bad.append(f'{sheet}R{r+1}图注(len={len(note_txt)})')
        # 四要素语义抽查: 线/符号/指针 注释特征词（至少 2 类）
        sem = sum(1 for k in ('线', '符号', '指针', '箭头', '曲线', '框')
                  if k in note_txt)
        if sem < 2:
            bad.append(f'{sheet}R{r+1}四要素(sem={sem})')
wb.close()
c_ok = n_img == 45 and not bad
check(f'C/D. {n_img} 幅: 图题黑体居中 + 图注左顶格四要素 + 锚点对齐', c_ok,
      f'异常={bad[:5]}')

fails = [r for r in results if not r[1]]
print(f'\n===== 全书图表终检（聚焦重绘版）: {"通过" if not fails else "未通过"} '
      f'({len(results)-len(fails)}/{len(results)}) =====')
sys.exit(0 if not fails else 1)
