# -*- coding: utf-8 -*-
"""《中医信息学》教材素材工作簿生成器
Book(GitBook 书稿) → Excel（目录结构 + 正文内容 + 真实挖掘数据 + embedding-code）。
产出：桌面 中医信息学.xlsx
"""
from __future__ import annotations
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from gitbook_to_excel import SW, parse_summary, sanitize, write_book  # noqa: E402

REPO = SCRIPTS.parent
BOOK = REPO / "kg" / "informatics" / "book"
CAT = REPO / "kg" / "informatics" / "tcm-informatics.json"
OUT = Path.home() / "Desktop" / "中医信息学.xlsx"

_T, _H, _B, _C = "标题", "标题", "正文", "代码"
F_T = Font(name="微软雅黑", size=13, bold=True, color="1F3864")
F_H = Font(name="微软雅黑", size=10.5, bold=True, color="2E5C8A")
F_B = Font(name="微软雅黑", size=10)
F_C = Font(name="Consolas", size=9.5)
FILL_H = PatternFill("solid", fgColor="DCE6F1")
FILL_T = PatternFill("solid", fgColor="EAF1F8")
THIN = Side(style="thin", color="BFBFBF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AL_L = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_toc(wb, chapters, used):
    ws = wb.create_sheet(sanitize("0目录细化", used))
    sw = SW(ws, [8, 30, 22, 62])
    sw.title("《中医信息学》目录结构（篇·章·节·目）", 4)
    sw.row(["级", "编号", "标题", "内容/说明"], [None] * 4,
           [F_H] * 4, [AL_C, AL_C, AL_L, AL_L], 20)
    part_no = ch_no = sec_no = 0
    for part, ct, cf in chapters:
        if part and not part.startswith("__"):
            part_no += 1
            sw.row([f"篇{part_no}", f"P{part_no}", part, f"（{ct}）"], [None] * 4,
                   [F_T, F_T, F_T, F_T], [AL_C, AL_C, AL_L, AL_L], 22)
    # 章节行
    for part, ct, cf in chapters:
        ch_no += 1
        sw.row([f"章", f"C{ch_no:02d}", ct, cf], [None] * 4,
               [F_H, F_H, F_H, F_B], [AL_C, AL_C, AL_L, AL_L], 18)
        p = BOOK / cf
        if p.exists():
            for line in p.read_text(encoding="utf-8").splitlines():
                m = re.match(r"^(#{2,3})\s+(.+)$", line.strip())
                if m:
                    lv = len(m.group(1)) - 1
                    sec_no += 1
                    sw.row([f"节{lv}", f"S{sec_no:03d}", m.group(2).strip(), ""], [None] * 4,
                           [F_B, F_B, F_B, F_B], [AL_C, AL_C, AL_L, AL_L], 16)
    return sw


def add_data_sheets(wb, cat, used):
    def new(name, widths):
        ws = wb.create_sheet(sanitize(name, used))
        return ws, SW(ws, widths)

    sw = new("信息标准", [16, 34, 20, 16, 40, 10])[1]
    sw.title("信息标准层（S）", 6)
    sw.row(["代码", "名称", "机构", "类别", "平台落地", "核实"], [None] * 6, [F_H] * 6,
           [AL_C] * 6, 20)
    for s in cat["standards"]:
        sw.row([s["code"], s["name"], s["org"], s["kind"], s["platform"], s["verify"]],
               [None] * 6, [F_B] * 6, [AL_C, AL_L, AL_C, AL_C, AL_L, AL_C], 18)

    sw = new("数据资产", [8, 26, 12, 8, 34, 40])[1]
    sw.title("数据资产层（A）", 6)
    sw.row(["编号", "名称", "规模", "单位", "路径", "说明"], [None] * 6, [F_H] * 6, [AL_C] * 6, 20)
    for a in cat["assets"]:
        sw.row([a["id"], a["name"], a["size"], a["unit"], a["path"], a["note"]],
               [None] * 6, [F_B] * 6, [AL_C, AL_L, AL_C, AL_C, AL_L, AL_L], 18)

    sw = new("挖掘算子", [8, 24, 40, 46])[1]
    sw.title("信息挖掘算子层（M）", 4)
    sw.row(["编号", "算子", "算法", "用途"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for o in cat["operators"]:
        sw.row([o["id"], o["name"], o["algo"], o["use"]], [None] * 4, [F_B] * 4,
               [AL_C, AL_L, AL_L, AL_L], 18)

    sw = new("频次分布", [22, 8, 22, 8])[1]
    sw.title("频次分布（M1）", 4)
    sw.row(["维度", "名次", "取值", "频次"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for key, items in cat["distributions"].items():
        for i, (name, n) in enumerate(items, 1):
            sw.row([key, i, name, n], [None] * 4, [F_B] * 4, [AL_C, AL_C, AL_L, AL_C], 16)

    sw = new("关联规则", [10, 40, 40, 10, 10, 10, 8])[1]
    sw.title("关联规则（M2，support/confidence/lift）", 7)
    sw.row(["类型", "前件 A", "后件 B", "support", "confidence", "lift", "n"], [None] * 7,
           [F_H] * 7, [AL_C] * 7, 20)
    for kind, rs in cat["rules"].items():
        for r in rs:
            sw.row([kind, r["left"], r["right"], r["support"], r["confidence"], r["lift"], r["n"]],
                   [None] * 7, [F_B] * 7, [AL_C, AL_L, AL_L, AL_C, AL_C, AL_C, AL_C], 16)

    sw = new("证候聚类", [10, 60])[1]
    sw.title("证候聚类（M3，重叠系数 ≥0.5）", 2)
    sw.row(["规模", "成员"], [None] * 2, [F_H] * 2, [AL_C] * 2, 20)
    for c in cat["clusters"]:
        sw.row([c["size"], "、".join(c["members"])], [None] * 2, [F_B] * 2, [AL_C, AL_L], 18)

    sw = new("编码覆盖", [10, 30, 16, 14])[1]
    sw.title(f"ICD-11 编码覆盖（M4）：{cat['icd_coverage']['mapped']}/{cat['icd_coverage']['total']}"
             f" = {cat['icd_coverage']['rate']:.1%}", 4)
    sw.row(["DSU", "疾病", "ICD-11", "ICD(旧)"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for r in cat["icd_coverage"]["unmapped_sample"]:
        sw.row([r["id"], r["disease"], r["icd11"] or "—", r["icd"] or "—"], [None] * 4,
               [F_B] * 4, [AL_C, AL_L, AL_C, AL_C], 16)

    sw = new("数据质量", [12, 26, 30, 26])[1]
    sw.title(f"数据质量门禁（M5）：{cat['stats']['data_quality_issues']} 项", 4)
    sw.row(["DSU", "字段", "取值/规则", "修复建议"], [None] * 4, [F_H] * 4, [AL_C] * 4, 20)
    for it in cat["data_quality"]["issues"]:
        sw.row([it["id"], it["field"], f"{it['value']} {it['rule']}".strip(), it["fix"]],
               [None] * 4, [F_B] * 4, [AL_C, AL_L, AL_L, AL_L], 16)


EMBED_CODE = [
 ("① 向量化 · 字符 bigram TF-IDF（kg/informatics/tcm_mining.py）",
  '''def bigrams(text):
    t = (text or "").lower(); toks = []
    for seg in re.findall(r"[\\u4e00-\\u9fff]+", t):
        toks += [seg] if len(seg) == 1 else [seg[i:i+2] for i in range(len(seg)-1)]
    for w in re.findall(r"[a-z0-9]{2,}", t):
        toks.append(w)
    return toks

class TfidfIndex:
    def _idf(self, term):
        return math.log((1 + self.n) / (1 + self.df.get(term, 0))) + 1.0
    def _vec(self, toks):
        tf = Counter(toks)
        v = {t: (c/len(toks)) * self._idf(t) for t, c in tf.items()} if toks else {}
        nrm = math.sqrt(sum(x*x for x in v.values())) or 1.0
        return {t: x/nrm for t, x in v.items()}'''),
 ("② 全文召回 · SQLite FTS5（scripts/tcm_embed.py）",
  '''CREATE VIRTUAL TABLE dsu_fts USING fts5(id UNINDEXED, text, tokenize='unicode61');
SELECT id, bm25(dsu_fts) FROM dsu_fts WHERE dsu_fts MATCH '失眠' ORDER BY 2 LIMIT 5;'''),
 ("③ 融合 · RRF（Reciprocal Rank Fusion，k=60）",
  '''def rrf(rank_lists, k=60):
    score = {}
    for rl in rank_lists:
        for i, doc in enumerate(rl):
            score[doc] = score.get(doc, 0.0) + 1.0 / (k + i + 1)
    return sorted(score, key=score.get, reverse=True)'''),
 ("④ ICD-11 桥接（scripts/icd11_client.py）",
  '''# 内网镜像三 header：Accept / Accept-Language:zh / API-Version:v2
import sqlite3
con = sqlite3.connect("tcmP/data/icd11_mms.db")     # 31,838 实体，2026-01
row = con.execute("SELECT id, code, title FROM entities WHERE id=?", (fid,)).fetchone()'''),
 ("⑤ 上下文工程装配（M7）",
  '''def build_context(q, k=5):
    hits = retrieve(q, k)                       # 向量 ∪ FTS5 ∪ RRF
    return "\\n".join(f"[{h.id}] {h.disease} / {h.syndrome}（证据等级 {h.evidence}）" for h in hits)
# 实测：Top-5 约 1.5K token，较全库 31K ↓≈95%'''),
 ("⑥ 数据质量门禁（M5）",
  '''EVIDENCE_ENUM = {"A级-多中心RCT","B级-单中心RCT","C级-队列研究",
                 "D级-病例对照","E级-专家共识","F级-经典理论"}
if ev and ev not in EVIDENCE_ENUM:
    issues.append(dict(id=uid, field="evidence_level", value=ev, fix="疑似 I级→E级"))'''),
]


def add_embed_sheet(wb, used):
    ws = wb.create_sheet(sanitize("embedding-code", used))
    sw = SW(ws, [30, 96])
    sw.title("embedding code：ICD-11 · FTS5 · TF-IDF · RRF · 上下文工程", 2)
    for title, code in EMBED_CODE:
        sw.row([title, ""], [None, None], [F_H, F_B], [AL_L, AL_L], 18)
        for ln in code.split("\n"):
            sw.row(["", ln], [None, None], [F_B, F_C], [AL_L, AL_L], 15 if ln.strip() else 12)


def main():
    cat = json.loads(CAT.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    meta = {"title": "中医信息学", "subtitle": "TCM Informatics",
            "author": "六者·中医医院AI网络教育平台 · tcmP",
            "note": "D07-S02 / AI-02 教材素材 · 与 kg/informatics 单一构建点同步"}
    write_book(wb, BOOK, meta, used)
    _, outline = parse_summary(BOOK / "SUMMARY.md")
    chapters = [(part, ct, cf) for part, chs in outline for ct, cf in chs]
    write_toc(wb, chapters, used)
    add_data_sheets(wb, cat, used)
    add_embed_sheet(wb, used)
    wb.save(OUT)
    print(f"[产出] {OUT}")
    print(f"[sheet] {len(wb.sheetnames)} 个：{', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()