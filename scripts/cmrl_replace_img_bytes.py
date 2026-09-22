# -*- coding: utf-8 -*-
"""cmrl_replace_img_bytes.py — 用最新 PNG 替换 Excel 内嵌图片字节（不插行）

用法: python scripts/cmrl_replace_img_bytes.py --ch N
图体行 = 图注行 - 1（布局: 图题行/图体行/图注行）
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from cmrl_merge import CHAP_META  # noqa: E402
from cmrl_insert_figures_v2 import FIGDIR, IMG_W, locate_figures_v2  # noqa: E402
from cmrl_update_xlsx import XLSX  # noqa: E402


def replace_images(ch: int):
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PImg
    prefix, sheet_name, title = CHAP_META[ch]
    figs = locate_figures_v2(ch)     # [(fig, para_idx, caption, note)]
    if not figs:
        print(f'[ch{ch:02d}] 未定位到图')
        return 0
    wb = load_workbook(XLSX)
    ws = wb[sheet_name]
    # 定位既有图注行: 扫描 C 列找 '图注：' 开头行（按顺序对应 figs）
    note_rows = [r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(r, 3).value or '').startswith('图注：')]
    if len(note_rows) != len(figs):
        print(f'[ch{ch:02d}] 图注行数({len(note_rows)}) != 图数({len(figs)}), 中止')
        wb.close()
        return 0
    # 清空旧图
    n_old = len(getattr(ws, '_images', None) or [])
    ws._images.clear()
    # 按图注行顺序替换: 图注行排序应与 figs 顺序一致（按 para_idx）
    figs_sorted = sorted(figs, key=lambda x: x[1])
    note_rows_sorted = sorted(note_rows)
    inserted = 0
    for (fig_name, para_idx, cap, note), note_row in zip(figs_sorted, note_rows_sorted):
        png = FIGDIR / f'ch{ch:02d}' / fig_name
        if not png.exists():
            print(f'  [MISS] {png}')
            continue
        body_row = note_row - 1     # 图体行 = 图注行上方一行
        w, h = PImg.open(png).size
        img = XLImage(str(png))
        img.width = IMG_W
        img.height = int(IMG_W * h / w)
        ws.add_image(img, f'C{body_row}')
        ws.row_dimensions[body_row].height = img.height * 0.75
        inserted += 1
        print(f'  [替换] {fig_name} → C{body_row} ({img.width}x{img.height}px)')
    wb.save(XLSX)
    print(f'[ch{ch:02d}] {sheet_name}: 替换 {inserted} 幅（旧 {n_old}）')
    return inserted


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--ch', type=int, required=True)
    args = ap.parse_args()
    replace_images(args.ch)
