# -*- coding: utf-8 -*-
"""constitution 模块测试套件（pytest）

运行：  python -m pytest -q          （根目录 pytest.ini 已把 testpaths 限定到本目录）
覆盖：  数据契约 / 生成引擎 / 逆转推理 / 性别参数 / 疾病易感预测 / 25型矩阵
"""
import os
import sys
import json

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from constitution import framework as F                     # noqa: E402
from constitution.model import ConstitutionModel            # noqa: E402

DATA = os.path.join(ROOT, 'constitution', 'data', 'wu_xing_25.json')


@pytest.fixture(scope='module')
def m():
    return ConstitutionModel()


# ── 数据契约 ────────────────────────────────────────────────
def test_json_contract_keys():
    kb = json.loads(open(DATA, encoding='utf-8').read())
    assert {'xing', 'types25', 'gender', 'hermes9', 'susceptibility'} <= set(kb)


def test_hermes9_nine_contiguous():
    kb = json.loads(open(DATA, encoding='utf-8').read())
    assert [h['no'] for h in kb['hermes9']] == list(range(1, 10))


def test_types25_shape(m):
    assert len(m.all_types()) == 25
    assert all(len(m.types[x]) == 5 for x in F.XING_ORDER)


def test_gender_compare_counts(m):
    cg = m.compare_gender()
    assert len(cg['same']) == 9 and len(cg['diff']) == 8
    assert all(r['dim'].startswith('★') for r in cg['diff'])


# ── 数值工具 ────────────────────────────────────────────────
@pytest.mark.parametrize('x,lv', [(0.10, 'L1'), (0.35, 'L2'), (0.50, 'L3'), (0.70, 'L4'), (0.85, 'L5')])
def test_level_of(x, lv):
    assert F.level_of(x) == lv


def test_gauss_peak_and_monotonic():
    assert abs(F.gauss(1.0) - 1.0) < 1e-9
    assert F.gauss(0.5) < F.gauss(0.8)


# ── 生成引擎：性别血气（LS65 气余血少）──────────────────────
def test_gender_qi_xue_law(m):
    pm = m.generate('木', '上角', 'male')['params']
    pf = m.generate('木', '上角', 'female')['params']
    assert pm['V_qi'] == 0.60 and pm['V_xue'] == 0.40
    assert pf['V_qi'] == 0.75 and pf['V_xue'] == 0.25
    assert pf['V_xue'] < pm['V_xue'] and pf['V_qi'] > pm['V_qi']


def test_hair_gender_dims(m):
    assert m.generate('木', '上角', 'male')['hair']['髯'] == '随血气盛衰'
    hf = m.generate('木', '上角', 'female')['hair']
    assert hf['髯'] == '0(无)' and hf['须'] == '0(无)' and hf['眉'] == '随血气盛衰'


def test_all_types_params_bounded(m):
    for t in m.all_types():
        for sex in ('male', 'female'):
            p = m.generate(t['xing'], t['name'], sex)['params']
            assert 0.0 <= p['V_qi'] <= 1.0 and 0.0 <= p['V_xue'] <= 1.0


# ── 形色相得/相克（LS64 五行生克）───────────────────────────
@pytest.mark.parametrize('xing,color,want', [
    ('木', '青', '相得'), ('木', '白', '色胜形'), ('木', '黄', '形胜色'), ('火', '黑', '色胜形')])
def test_form_color_relation(m, xing, color, want):
    assert want in m._susceptibility(xing, color)['形色关系']


# ── 逆转推理 ────────────────────────────────────────────────
def test_infer_water_taiyu_with_position(m):
    r = m.infer({'color': '黑', 'yin': '羽', 'jing': 'BL', 'pos': '右足太阳之上',
                 'shape': ['大头', '大腹']})
    assert r['main']['code'] == '水形·太羽'
    assert r['ranked'][0]['score'] > r['ranked'][1]['score']


def test_infer_main_type_bonus(m):
    r = m.infer({'color': '青', 'yin': '角', 'jing': 'LR', 'shape': ['小头', '长面', '大肩背']})
    assert r['main']['code'] == '木形·上角'


# ── 疾病易感预测 ────────────────────────────────────────────
def test_susceptibility_fire(m):
    s = m.predict_susceptibility('火', None, 'male')
    assert s['zang'] == '心'
    assert any('冠心病' in d for d in s['diseases'])
    assert 'Type A' in s['modern'] or 'A型' in s['modern']
    assert s['risk_level'] == '基线' and len(s['nian_ji']) == 7


def test_susceptibility_sex_specific(m):
    sm = m.predict_susceptibility('火', None, 'male')
    sf = m.predict_susceptibility('火', None, 'female')
    assert sm['sex_specific'].startswith('男') and sf['sex_specific'].startswith('女')
    assert sm['sex_specific'] != sf['sex_specific']


def test_susceptibility_risk_float_on_clash(m):
    s = m.predict_susceptibility('木', None, 'male', color_se='白')   # 金克木
    assert '上浮' in s['risk_level']


# ── 路线图 ──────────────────────────────────────────────────
def test_roadmap_phases(m):
    rp = m.roadmap()
    assert len(rp) == 8 and rp[0]['phase'].startswith('P1') and rp[-1]['phase'].startswith('P8')


# ── 交付物 docx（若已生成则回读）────────────────────────────
DOCX = r'C:\Users\DELL\Desktop\yy25.docx'


@pytest.mark.skipif(not os.path.exists(DOCX), reason='yy25.docx 未生成')
def test_yy25_docx_artifact():
    from docx import Document
    d = Document(DOCX)
    cell = [c.text for t in d.tables for r in t.rows for c in r.cells]
    full = '\n'.join([p.text for p in d.paragraphs] + cell)
    assert os.path.getsize(DOCX) > 50000
    assert len(d.inline_shapes) == 2 and len(d.tables) >= 7
    for kw in ['模型架构图', '落地路线图', '男女参数异同', '模型预测疾病易感性',
               '五五二十五人', '表面组学', '行为心脏病学', '有余于气', '不足于血']:
        assert kw in full


# ══════════════════════════════════════════════════════════════
# V2 术数层：双层进制 / 60甲子 / 男8女7 / 左右镜像 / 双目标
# ══════════════════════════════════════════════════════════════

def test_dual_radix(m):
    d = m.dual_radix()
    assert d['radix_low']['base'] == 2 and d['radix_high']['base'] == 5
    assert '10' in d['composite'] and '60' in d['jiazi_cycle']
    assert len(d['tiangan']) == 10 and len(d['dizhi']) == 12


def test_jiazi60_cycle():
    assert len(F.JIAZI60) == 60
    assert F.JIAZI60[0]['gz'] == '甲子' and F.JIAZI60[59]['gz'] == '癸亥'
    assert len({j['gz'] for j in F.JIAZI60}) == 60          # 60甲子互异
    assert F.JIAZI60[0]['tg'] == '甲' and F.JIAZI60[0]['dz'] == '子'
    assert F.JIAZI60[0]['tg_wx'] == '木' and F.JIAZI60[0]['shengxiao'] == '鼠'


def test_jiazi_by_year(m):
    assert m.jiazi_by_year(1984)['gz'] == '甲子'
    assert m.jiazi_by_year(2024)['gz'] == '甲辰'


def test_jiazi_lifetime_relation(m):
    lt = m.jiazi_lifetime(0, '木')          # 甲子 → 甲属木，与木形同气
    assert lt['ming_wuxing'] == '木' and lt['xing_relation'] == '同'
    lt2 = m.jiazi_lifetime(0, '金')         # 木 vs 金 → 被克(金克木)
    assert lt2['xing_relation'] == '被克'


@pytest.mark.parametrize('a,b,want', [
    ('木', '木', '同'), ('木', '火', '生'), ('火', '木', '被生'),
    ('木', '土', '克'), ('土', '木', '被克')])
def test_wuxing_relation(a, b, want):
    assert F.wuxing_relation(a, b) == want


def test_development_female_male(m):
    f35 = m.development('female', 35)
    assert f35['base'] == 7 and f35['stage'] == '始衰' and f35['tiangui'] == '已至'
    assert m.development('female', 14)['tiangui'] == '已至'
    assert m.development('female', 13)['tiangui'] == '未至'
    assert m.development('female', 49)['tiangui'] == '已竭'
    m40 = m.development('male', 40)
    assert m40['base'] == 8 and m40['stage'] == '始衰'


def test_mirror_rule(m):
    assert m.mirror('male')['first'] == '左' and m.mirror('female')['first'] == '右'
    assert F.side_yinyang('左', 'male') == '阳'
    assert F.side_yinyang('左', 'female') == '阴'          # 男女镜像
    assert F.side_yinyang('右', 'female') == '阳'
    assert m.subtype_side_sex('太角(左足少阳之上)', 'male') == '阳'
    assert m.subtype_side_sex('太角(左足少阳之上)', 'female') == '阴'


def test_health_state_target1(m):
    h = m.health_state('火', 'female', 2, 28)              # 形色相得+命局同气+巅峰
    assert 0.0 <= h['health_score'] <= 1.0
    assert h['state'] == '平人' and h['dev_stage'] == '巅峰'
    assert h['mirror'] == '右为阳'
    # 相克 + 衰期 → 下降
    bad = m.health_state('木', 'male', None, 64, )['health_score']
    assert bad < h['health_score']


def test_syndrome_susceptibility_target2(m):
    s = m.syndrome_susceptibility('火', 'female', 2, 28)   # 命局火 → 心火亢盛证↑
    assert s['main_zheng'] == '心火亢盛证' and s['bing_xing'] == '热/瘀'
    vals = [r['susceptibility'] for r in s['ranked']]
    assert vals == sorted(vals, reverse=True)              # 降序
    assert all(0.0 <= v <= 1.0 for v in vals)
    # 女以血为用 → 肝血虚证 出现在排序且不低于 L3 基线
    assert any(r['zheng'] == '肝血虚证' for r in s['ranked'])
    # 老年 → 肾虚类加权
    s2 = m.syndrome_susceptibility('水', 'male', None, 60)
    assert any('肾' in r['zheng'] for r in s2['ranked'])

